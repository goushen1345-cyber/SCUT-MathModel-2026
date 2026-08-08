"""Q2: Minimum suppliers, economical ordering and low-loss transport."""
from __future__ import annotations

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook

from config import ATTACH_A, ATTACH_B, PLAN_WEEKS, PROCESSED_DIR, TABLE_DIR, WEEKLY_CAPACITY
from optimization import (
    PlanConfig,
    assign_transport,
    build_loss_matrix,
    evaluate_plan,
    mix_weekly_raw_needs,
    select_min_suppliers,
    solve_order_plan,
    weekly_raw_demand,
)
from q1_supplier_importance import rank_suppliers
from utils import build_supplier_features, ensure_dirs, load_supplier_data, load_transporter_data, setup_logging, setup_matplotlib


def write_attachment_a(order_df: pd.DataFrame, sheet_name: str):
    wb = load_workbook(ATTACH_A)
    ws = wb[sheet_name]
    week_cols = [f"week_{i:02d}" for i in range(1, PLAN_WEEKS + 1)]
    plan = order_df.set_index("supplier_id")
    for r in range(6, 408):
        sid = ws.cell(r, 1).value
        if sid in plan.index:
            for c, wk in enumerate(range(2, 26)):
                val = plan.loc[sid, f"week_{wk-1:02d}"]
                ws.cell(r, c + 2, value=round(val, 2) if val and val > 0.01 else None)
    out = PROCESSED_DIR / f"附件A_{sheet_name}.xlsx"
    wb.save(out)
    return out


def write_attachment_b(transport: pd.DataFrame, sheet_name: str):
    wb = load_workbook(ATTACH_B)
    ws = wb[sheet_name]
    for _, row in transport.iterrows():
        sid, week, tname, vol = row["supplier_id"], int(row["week"]), row["transporter"], row["volume"]
        if vol <= 0.01:
            continue
        # find supplier row
        srow = None
        for r in range(6, 408):
            if ws.cell(r, 1).value == sid:
                srow = r
                break
        if srow is None:
            continue
        t_idx = int(tname[1:]) - 1
        col = 2 + (week - 1) * 8 + t_idx
        ws.cell(srow, col, value=round(vol, 2))
    out = PROCESSED_DIR / f"附件B_{sheet_name}.xlsx"
    wb.save(out)
    return out


def main():
    ensure_dirs()
    logger = setup_logging("q2_order_transport")
    setup_matplotlib()
    orders, supplies = load_supplier_data()
    trans = load_transporter_data()
    features = build_supplier_features(orders, supplies)
    ranked, _ = rank_suppliers(features)

    cfg = PlanConfig(objective="cost")
    needs = mix_weekly_raw_needs(cfg)
    total_raw = sum(needs.values())
    supplier_ids, min_n = select_min_suppliers(ranked, total_raw)

    order_plan = solve_order_plan(features, supplier_ids, cfg)
    loss_mat = build_loss_matrix(trans)
    transport = assign_transport(order_plan, loss_mat)
    metrics = evaluate_plan(order_plan, transport, cfg)

    order_plan.to_csv(TABLE_DIR / "q2_order_plan.csv", index=False, encoding="utf-8-sig")
    transport.to_csv(TABLE_DIR / "q2_transport_plan.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame([{"min_suppliers": min_n, **metrics}]).to_csv(
        TABLE_DIR / "q2_summary.csv", index=False, encoding="utf-8-sig"
    )

    write_attachment_a(order_plan, "问题2的订购方案结果")
    write_attachment_b(transport, "问题2的转运方案结果")

    # figures
    week_cols = [f"week_{i:02d}" for i in range(1, PLAN_WEEKS + 1)]
    weekly = order_plan[week_cols].sum()
    fig, ax = plt.subplots()
    ax.bar(range(1, PLAN_WEEKS + 1), weekly.values, color="#27AE60")
    ax.axhline(total_raw, color="r", ls="--", label=f"需求={total_raw:.0f}")
    ax.set_xlabel("周")
    ax.set_ylabel("订货量 (m³)")
    ax.set_title("Q2 24周订货量")
    ax.legend()
    fig.tight_layout()
    fig.savefig("output/figures/q2_weekly_orders.png")
    plt.close(fig)

    recv_loss = []
    for w in range(PLAN_WEEKS):
        sub = transport[transport["week"] == w + 1]
        vol = sub["volume"].sum()
        # weighted loss
        loss_w = loss_mat[:, w]
        wloss = 0.0
        for ti in range(8):
            v = sub[sub["transporter"] == f"T{ti+1}"]["volume"].sum()
            wloss += v * loss_w[ti]
        recv_loss.append(wloss / max(vol, 1))
    fig, ax = plt.subplots()
    ax.plot(range(1, PLAN_WEEKS + 1), recv_loss, "o-", color="#8E44AD")
    ax.set_xlabel("周")
    ax.set_ylabel("加权平均损耗率 (%)")
    ax.set_title("Q2 转运方案损耗效果")
    fig.tight_layout()
    fig.savefig("output/figures/q2_transport_loss.png")
    plt.close(fig)

    logger.info("Q2: min_suppliers=%d, total_order=%.0f", min_n, metrics["total_order_volume"])
    return {"min_suppliers": min_n, "metrics": metrics, "supplier_ids": supplier_ids}


if __name__ == "__main__":
    main()
