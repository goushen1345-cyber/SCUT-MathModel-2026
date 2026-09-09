from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
FINAL = ROOT / "paper" / "final" / "results"

ITEMS = [
    {
        "name": "result1_1.xlsx",
        "source": ROOT / "questions" / "Q1" / "outputs" / "Q1-DRAFT-v2" / "result1_1_Q1-DRAFT-v2.xlsx",
        "prior": ROOT / "questions" / "Q1" / "outputs" / "Q1-DRAFT-v2" / "workbook_validation.json",
        "expected_nonzero": 738,
    },
    {
        "name": "result1_2.xlsx",
        "source": ROOT / "questions" / "Q1" / "outputs" / "Q1-DRAFT-v1" / "result1_2_Q1-DRAFT-v1.xlsx",
        "prior": ROOT / "questions" / "Q1" / "outputs" / "Q1-DRAFT-v1" / "workbook_validation.json",
        "expected_nonzero": 662,
    },
    {
        "name": "result2.xlsx",
        "source": ROOT / "questions" / "Q2" / "outputs" / "Q2-DRAFT-v1" / "result2_Q2-DRAFT-v1.xlsx",
        "prior": ROOT / "questions" / "Q2" / "outputs" / "Q2-DRAFT-v1" / "workbook_validation.json",
        "expected_nonzero": 666,
    },
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def inspect(path: Path) -> dict:
    wb = load_workbook(path, read_only=True, data_only=False)
    names = wb.sheetnames
    nonzero = 0
    formula_errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, (int, float)) and not isinstance(value, bool) and abs(value) > 1e-9:
                    nonzero += 1
                if isinstance(value, str) and value.startswith("#") and value in {
                    "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"
                }:
                    formula_errors.append(f"{ws.title}!{cell.coordinate}:{value}")
    wb.close()
    return {"sheet_names": names, "nonzero_numeric_cells": nonzero, "formula_errors": formula_errors}


def main() -> None:
    FINAL.mkdir(parents=True, exist_ok=True)
    results = []
    for item in ITEMS:
        target = FINAL / item["name"]
        shutil.copy2(item["source"], target)
        prior = json.loads(item["prior"].read_text(encoding="utf-8"))
        audit = inspect(target)
        copy_hash = sha256(target)
        source_hash = sha256(item["source"])
        record = {
            "file": item["name"],
            "source": item["source"].relative_to(ROOT).as_posix(),
            "source_sha256": source_hash,
            "copy_sha256": copy_hash,
            "byte_identical_to_frozen_source": source_hash == copy_hash,
            "sheet_names": audit["sheet_names"],
            "seven_year_sheets": audit["sheet_names"] == [str(y) for y in range(2024, 2031)],
            "nonzero_numeric_cells": audit["nonzero_numeric_cells"],
            "expected_nonzero_from_prior_independent_check": item["expected_nonzero"],
            "nonzero_matches_prior": audit["nonzero_numeric_cells"] == item["expected_nonzero"],
            "formula_errors": audit["formula_errors"],
            "prior_independent_validation_passed": bool(prior.get("passed")),
        }
        record["passed"] = all([
            record["byte_identical_to_frozen_source"],
            record["seven_year_sheets"],
            record["nonzero_matches_prior"],
            not record["formula_errors"],
            record["prior_independent_validation_passed"],
        ])
        results.append(record)
    report = {
        "protocol": "PAPER-FINAL-v1-workbook-review",
        "method": "byte-for-byte frozen-source copy plus independent openpyxl readback; no solver or workbook builder imported",
        "files": results,
        "passed": all(x["passed"] for x in results),
        "human_verification": "pending",
    }
    (FINAL / "workbook_validation.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if not report["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
