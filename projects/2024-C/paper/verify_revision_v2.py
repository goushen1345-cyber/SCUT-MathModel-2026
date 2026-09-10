"""Read-only verification of existing evidence and newly rendered paper/figures.

No research modules, scenario generators, optimizers or workbook writers are used.
Only final_v2/validation.json is written.
"""
import csv
import hashlib
import json
import re
import zipfile
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from PIL import Image
from pypdf import PdfReader

PAPER = Path(__file__).resolve().parent
PROJECT = PAPER.parent
FIG = PAPER / 'figures/overall_v2'
OUT = PAPER / 'final_v2'


def read_json(path):
    return json.loads(path.read_text(encoding='utf-8-sig'))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def rows(path):
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def main():
    checks = []

    def check(name, passed, details):
        checks.append({'check': name, 'passed': bool(passed), 'details': details})

    state = read_json(PROJECT / 'state/state.json')
    manifests = dict(state['data_snapshots'])
    for q, info in state['questions'].items():
        manifests[q] = info['snapshots'][info['frozen_version']]
    counts, mismatches = {}, []
    for scope, files in manifests.items():
        counts[scope] = len(files)
        for rel, expected in files.items():
            p = PROJECT / rel
            if not p.is_file() or sha(p) != expected:
                mismatches.append(rel)
    check('frozen_sources_unchanged', not mismatches and sum(counts.values()) == 288,
          {'counts': counts, 'mismatches': mismatches})

    old_hashes = {
        'final/2024-C-paper.docx': '622c9bda9177ce770ed21723c0bbfb8509b419bdfa58fc2e9b1bbbc632b60dda',
        'final/2024-C-paper.pdf': '209a0e300d419498699e08fc1cac6acd24d0b0a03b092613d3dce2ec42d6c920',
        'manuscript_v1.md': read_json(PAPER / 'corrections_v2.json')['source_sha256'],
    }
    old_results = {p: sha(PAPER / p) == h for p, h in old_hashes.items()}
    check('original_paper_preserved', all(old_results.values()), old_results)
    wb = read_json(PAPER / 'final/results/workbook_validation.json')['files']
    workbooks = {r['file']: sha(PAPER / 'final/results' / r['file']) ==
                 sha(PROJECT / r['source']) == r['source_sha256'] for r in wb}
    check('three_result_workbooks_unchanged', len(workbooks) == 3 and all(workbooks.values()), workbooks)
    coverage = rows(PROJECT / 'data/frozen/G1-v1-base/coverage.csv')
    check('unknown_sales_preserved', len(coverage) == 59 and
          all(r['normal_sales_limit_jin'] == '' and r['demand_status'] == 'unknown_not_imputed' for r in coverage),
          {'rows': len(coverage), 'all_sales_unknown': all(r['normal_sales_limit_jin'] == '' for r in coverage)})

    holdout = rows(PROJECT / 'questions/Q3/deep_runs/holdout_summary.csv')[0]
    d = {k: Decimal(holdout[k]) for k in ['mean', 'R0_mean', 'RS_mean', 'RSC_mean']}
    rs, rsc = d['RS_mean'] - d['R0_mean'], d['RSC_mean'] - d['R0_mean']
    correction = read_json(PAPER / 'corrections_v2.json')
    check('holdout_differences_from_unrounded_values', rs == Decimal('-1916.070086') and
          rsc == Decimal('-605.608423') and str(rs) == correction['table6']['raw_RS_minus_R0_yuan'] and
          str(rsc) == correction['table6']['raw_RSC_minus_R0_yuan'],
          {'RS_minus_R0_yuan': str(rs), 'RSC_minus_R0_yuan': str(rsc), 'sample': '500 shared base scenarios; 3 worlds'})
    diag = rows(PROJECT / 'questions/Q1/deep_runs/bean_diagnostic.csv')[0]
    prior_gap = Decimal(diag['main_profit_yuan']) - Decimal(diag['profit_yuan'])
    check('historical_bean_baseline', prior_gap == Decimal('1683085.625'),
          {'old_main_yuan': diag['main_profit_yuan'], 'old_strict_yuan': diag['profit_yuan'],
           'old_difference_yuan': str(prior_gap), 'diagnostic_rerun': False})

    doc = Document(OUT / '2024-C-paper-v2.docx')
    text = '\n'.join([p.text for p in doc.paragraphs] +
                     [c.text for t in doc.tables for r in t.rows for c in r.cells])
    md = (PAPER / 'manuscript_v2.md').read_text(encoding='utf-8')
    pdf = PdfReader(OUT / '2024-C-paper-v2.pdf')
    pdf_text = '\n'.join(p.extract_text() or '' for p in pdf.pages)
    compact = lambda x: re.sub(r'\s+', '', x)
    anchors = ['5%和10%', '年度代理总量不变', '交叉价格响应', '额外关系乘子',
               '−1916.07', '−605.61', '成本和蔬菜价格', '偏离了原方案',
               '训练样本的潜在变量层', '最终参数层秩相关目标', '补强前', '严格诊断未重新求解']
    anchor_results = {s: {'md': compact(s) in compact(md), 'docx': compact(s) in compact(text),
                          'pdf': compact(s) in compact(pdf_text)} for s in anchors}
    check('six_corrections_present_in_all_formats', all(all(v.values()) for v in anchor_results.values()), anchor_results)
    check('manuscript_hash_registered', sha(PAPER / 'manuscript_v2.md') == correction['output_sha256'],
          sha(PAPER / 'manuscript_v2.md'))
    check('paper_structure', len(doc.inline_shapes) == 9 and len(doc.tables) == 8 and len(pdf.pages) == 18 and
          md.count('[[EQ:') == 11, {'figures': len(doc.inline_shapes), 'tables': len(doc.tables),
                                  'pages': len(pdf.pages), 'numbered_equations': md.count('[[EQ:'),
                                  'math_objects': len(doc._element.xpath('//m:oMath'))})
    check('corrected_table6_numbers', all(v in text for v in ['6727.36', '6727.17', '6727.30', '6727.28', '−1916.07', '−605.61']),
          correction['table6']['rows'])
    w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    colors = sorted({x.get(w+'val') for x in doc._element.iter(w+'color')})
    fills = sorted({x.get(w+'fill') for x in doc._element.iter(w+'shd') if x.get(w+'fill')})
    check('black_text_white_tables', set(colors) <= {'000000', 'auto'} and set(fills) <= {'FFFFFF', 'auto'},
          {'colors': colors, 'fills': fills})
    with zipfile.ZipFile(OUT / '2024-C-paper-v2.docx') as z:
        svg_parts = [n for n in z.namelist() if n.startswith('word/media/') and n.endswith('.svg')]
        check('native_vector_figures_embedded', len(svg_parts) == 2 and
              z.read('word/media/figure_vector_0.svg') == (FIG / 'render/overall_framework.svg').read_bytes() and
              z.read('word/media/figure_vector_7.svg') == (FIG / 'render/q3_relationship.svg').read_bytes(), svg_parts)

    spec = read_json(FIG / 'structure.json')
    forbidden = ['agent', 'skill', 'reviewer', 'freeze', 'state', 'wait', 'checkpoint', 'prompt',
                 '用户确认', '内部文件路径', 'M0', 'M1', 'B0', 'N10']
    public = '\n'.join('\n'.join(n['lines']) for n in spec['nodes'])
    check('figure_structure_and_public_labels', len(spec['nodes']) == 19 and len(spec['edges']) == 18 and
          all(2 <= len(n['lines']) <= 4 for n in spec['nodes']) and
          not any(s.lower() in public.lower() for s in forbidden),
          {'nodes': len(spec['nodes']), 'edge_records': len(spec['edges']), 'shared_group_members': ['UV', 'UA', 'UT']})
    source_hashes = read_json(FIG / 'source_manifest.json')['sources']
    stale = [p for p, h in source_hashes.items() if sha(PROJECT / p) != h]
    check('figure_sources_traceable', not stale, {'sources': len(source_hashes), 'stale': stale})
    render = read_json(FIG / 'render/render_manifest.json')
    render_bad = [p for p, h in render['artifacts'].items() if sha(FIG / 'render' / p) != h]
    check('render_manifest_current', not render_bad and render['structure_sha256'] == sha(FIG / 'structure.json'), render_bad)
    rep = {}
    for stem in ['overall_framework', 'q3_relationship']:
        a, b = FIG / 'render', FIG / 'reproduction'
        ai, bi = Image.open(a / (stem+'.png')).convert('RGB'), Image.open(b / (stem+'.png')).convert('RGB')
        svg = ET.parse(a / (stem+'.svg')).getroot()
        rep[stem] = {
            'tex_identical': (a / (stem+'.tex')).read_bytes() == (b / (stem+'.tex')).read_bytes(),
            'svg_identical': (a / (stem+'.svg')).read_bytes() == (b / (stem+'.svg')).read_bytes(),
            'png_pixels_identical': ai.size == bi.size and ai.tobytes() == bi.tobytes(),
            'pdf_one_page': len(PdfReader(a / (stem+'.pdf')).pages) == 1,
            'svg_has_vector_paths': bool(svg.findall('.//{http://www.w3.org/2000/svg}path')),
            'svg_has_no_raster_images': not svg.findall('.//{http://www.w3.org/2000/svg}image'),
        }
    check('two_figures_reproduced_as_vectors', all(all(v.values()) for v in rep.values()), rep)
    pages = sorted((OUT / 'qa_pages_final').glob('page-*.png'))
    check('all_final_pdf_pages_rendered', len(pages) == len(pdf.pages), {'pages': len(pages), 'dpi': 150})
    artifact_paths = [OUT/'2024-C-paper-v2.docx', OUT/'2024-C-paper-v2.pdf', PAPER/'manuscript_v2.md',
                      FIG/'structure.json', PAPER/'build_paper_v2.py', PAPER/'verify_revision_v2.py']
    result = {
        'version': 'PAPER-FINAL-v2-targeted-verification', 'date_utc': datetime.now(timezone.utc).isoformat(),
        'method': 'Static sources, hashes, Decimal arithmetic, document readback and figure reproduction comparison; no research code imported',
        'state_revision_at_check': state['revision'], 'checks': checks,
        'passed': all(c['passed'] for c in checks), 'new_experiments_run': 0,
        'visual_review': {'pages': '01—18', 'source': 'qa_pages_final, actual PDF at 150 dpi',
                          'method': 'Two read-only AI subtasks viewed all 18 pages, root viewed corrected figure/math pages',
                          'resolved': ['figure title glyph overlap', 'table5 separated header', 'short isolated paragraph page', 'raw formula notation'],
                          'remaining_minor': ['page6 whitespace below table2', 'unchanged old figure4 small annotations/no axis ticks'],
                          'human_review_claimed': False},
        'limits': ['Targeted correction and rendering audit, not a new full scientific validation',
                   'Protocol/implementation deviation remains disclosed, not retrospectively cured',
                   'Final economic-parameter correlation targets remain unverified',
                   'AI agreement is not proof or participant approval'],
        'artifact_sha256': {str(p.relative_to(PROJECT)).replace('\\', '/'): sha(p) for p in artifact_paths},
    }
    (OUT / 'validation.json').write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    print(json.dumps({'passed': result['passed'], 'checks': len(checks),
                      'failed': [c['check'] for c in checks if not c['passed']]}))
    if not result['passed']:
        raise SystemExit(1)


if __name__ == '__main__':
    main()
