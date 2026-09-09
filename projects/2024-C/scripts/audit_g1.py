"""Read-only source audit. Writes evidence summaries, never cleaned input tables."""
from pathlib import Path
from collections import Counter, defaultdict
from datetime import datetime
import hashlib
import json
import re
import warnings
import zipfile
import xml.etree.ElementTree as ET
from decimal import Decimal
import openpyxl

PROJECT = Path(__file__).resolve().parents[1]
SOURCE = PROJECT.parents[1] / '培训作业' / '第四次培训作业'
OUT = PROJECT / 'audit' / 'G1-v1'

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def diagnostic_value(ws, row, col):
    """Merged-anchor lookup for auditing only; not an approved fill operation."""
    c = ws.cell(row, col)
    if c.value is not None:
        return c.value
    for area in ws.merged_cells.ranges:
        if area.min_row <= row <= area.max_row and area.min_col <= col <= area.max_col:
            return ws.cell(area.min_row, area.min_col).value
    return None

def norm(x):
    return x.strip() if isinstance(x, str) else x

def dupes(keys):
    return {str(k): v for k, v in Counter(keys).items() if v > 1}

def rows_with_id(ws, col):
    return [r for r in range(2, ws.max_row+1) if isinstance(ws.cell(r, col).value, (int,float))]

def main():
    paths = sorted(p for p in SOURCE.rglob('*') if p.is_file())
    before = {str(p.relative_to(SOURCE)):sha(p) for p in paths}
    out = {'scope':'Read-only structural and quality audit; diagnostic normalization is not an adopted data treatment.',
           'source_directory':str(SOURCE), 'source_hashes_before':before,
           'openpyxl_version':openpyxl.__version__, 'books':{}}
    books = {}
    for p in paths:
        if p.suffix != '.xlsx':
            continue
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter('always')
            w = openpyxl.load_workbook(p, data_only=False)
        books[p.name] = w
        meta = {'sheets':{},'load_warnings':sorted(set(str(x.message) for x in caught))}
        for ws in w:
            populated = [c for row in ws for c in row if c.value is not None]
            meta['sheets'][ws.title] = {
                'dimensions':[ws.max_row,ws.max_column],
                'merged_ranges':sorted(str(x) for x in ws.merged_cells.ranges),
                'nonempty_cells':len(populated),
                'formulas':[c.coordinate for c in populated if c.data_type=='f'],
                'errors':[{ 'cell':c.coordinate, 'value':c.value} for c in populated if c.data_type=='e'],
                'whitespace_cells':[{'cell':c.coordinate,'value':c.value} for c in populated if isinstance(c.value,str) and c.value!=c.value.strip()],
                'comments':[{'cell':c.coordinate,'text':c.comment.text} for c in populated if c.comment],
                'hidden_rows':[r for r,d in ws.row_dimensions.items() if d.hidden],
                'hidden_columns':[c for c,d in ws.column_dimensions.items() if d.hidden],
            }
        out['books'][str(p.relative_to(SOURCE))] = meta

    land_ws,crop_ws = books['附件1.xlsx'].worksheets
    plant_ws,param_ws = books['附件2.xlsx'].worksheets
    land_rows = [r for r in range(2,land_ws.max_row+1) if re.fullmatch(r'[A-F]\d+',str(land_ws.cell(r,1).value))]
    lands = {land_ws.cell(r,1).value:{'row':r,'type':norm(land_ws.cell(r,2).value),'area':land_ws.cell(r,3).value} for r in land_rows}
    land_groups = defaultdict(list)
    for k,v in lands.items(): land_groups[v['type']].append((k,v['area']))
    crop_rows=rows_with_id(crop_ws,1)
    crops={crop_ws.cell(r,1).value:{'row':r,'name':norm(crop_ws.cell(r,2).value),'type':norm(crop_ws.cell(r,3).value)} for r in crop_rows}
    out['land']={'count':len(land_rows),'duplicate_ids':dupes([land_ws.cell(r,1).value for r in land_rows]),
                 'by_type':{k:{'count':len(v),'area_sum':sum(x[1] for x in v),'min_area':min(x[1] for x in v),'max_area':max(x[1] for x in v)} for k,v in land_groups.items()},
                 'nonpositive_area':[k for k,v in lands.items() if not isinstance(v['area'],(int,float)) or v['area']<=0]}
    out['crops']={'count':len(crop_rows),'duplicate_ids':dupes([crop_ws.cell(r,1).value for r in crop_rows]),'by_type':dict(Counter(v['type'] for v in crops.values()))}

    pr = rows_with_id(plant_ws,2)
    planting=[]
    problems=[]
    blank_examples=[]
    for r in pr:
        raw=[plant_ws.cell(r,c).value for c in range(1,7)]
        vals=[diagnostic_value(plant_ws,r,c) for c in range(1,7)]
        if raw[0] is None: blank_examples.append({'cell':f'A{r}','anchor_value_for_diagnosis':vals[0]})
        land,cid,name,ctype,area,season=map(norm,vals)
        item={'row':r,'land':land,'crop_id':cid,'name':name,'type':ctype,'area':area,'season':season}
        planting.append(item)
        if any(v is None for v in vals): problems.append({'row':r,'kind':'missing_after_merge_lookup','values':vals})
        if land not in lands: problems.append({'row':r,'kind':'unknown_land','value':land})
        if cid not in crops: problems.append({'row':r,'kind':'unknown_crop','value':cid})
        elif name!=crops[cid]['name'] or ctype!=crops[cid]['type']: problems.append({'row':r,'kind':'crop_label_mismatch'})
        if not isinstance(area,(int,float)) or area<=0: problems.append({'row':r,'kind':'nonpositive_area','value':area})
    area_groups=defaultdict(float)
    crop_count_groups=Counter()
    for p in planting:
        area_groups[(p['land'],p['season'])]+=p['area']
        crop_count_groups[(p['land'],p['season'])]+=1
    out['planting']={'record_count':len(pr),'blank_land_cells':blank_examples,'problems':problems,
        'duplicate_land_crop_season_keys':dupes([(p['land'],p['crop_id'],p['season']) for p in planting]),
        'season_counts':dict(Counter(p['season'] for p in planting)),
        'area_sum_by_season':{s:sum(p['area'] for p in planting if p['season']==s) for s in sorted(set(p['season'] for p in planting))},
        'land_season_area_mismatches':[{'land':k[0],'season':k[1],'planted':v,'land_area':lands[k[0]]['area']} for k,v in area_groups.items() if abs(v-lands[k[0]]['area'])>1e-8],
        'mixed_land_seasons':[{'land':k[0],'season':k[1],'crop_count':v} for k,v in crop_count_groups.items() if v>1],
        'min_planted_area':min(p['area'] for p in planting),'max_planted_area':max(p['area'] for p in planting)}

    rr=rows_with_id(param_ws,1)
    params={}
    pkeys=[]
    price_fail=[]
    nonpositive=[]
    label_fail=[]
    price_by_crop=defaultdict(set)
    for r in rr:
        a=[norm(param_ws.cell(r,c).value) for c in range(1,9)]
        _,cid,name,lt,season,yield_,cost,price=a
        key=(cid,lt,season);pkeys.append(key);params[key]={'row':r,'yield':yield_,'cost':cost,'price':price}
        if cid not in crops or name!=crops[cid]['name']:label_fail.append({'row':r,'values':a})
        if not all(isinstance(v,(int,float)) and v>0 for v in (yield_,cost)):nonpositive.append({'row':r,'yield':yield_,'cost':cost})
        m=re.fullmatch(r'(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)',str(price))
        if not m or float(m[1])>float(m[2]) or float(m[1])<=0:price_fail.append({'row':r,'price':price})
        price_by_crop[cid].add(price)
    matches=Counter();missing=[];note_inheritance=[]
    for p in planting:
        lt=lands[p['land']]['type'];key=(p['crop_id'],lt,p['season'])
        if key in params:matches['direct']+=1
        elif lt=='智慧大棚' and p['season']=='第一季' and (p['crop_id'],'普通大棚','第一季') in params:
            matches['explicit_note_equivalence']+=1
            note_inheritance.append({'planting_row':p['row'],'source_parameter_row':params[(p['crop_id'],'普通大棚','第一季')]['row'],'note':'2023年统计的相关数据!B111'})
        else:matches['unmatched']+=1;missing.append({'planting_row':p['row'],'key':key})
    out['parameters']={'record_count':len(rr),'duplicate_crop_land_season_keys':dupes(pkeys),'label_mismatches':label_fail,
        'nonpositive_yield_cost':nonpositive,'price_parse_failures':price_fail,
        'multiple_price_ranges_per_crop':{str(k):sorted(v) for k,v in price_by_crop.items() if len(v)>1},
        'observed_planting_match_counts':dict(matches),'unmatched_planting_records':missing,'explicit_note_equivalences':note_inheritance,
        'yield_range':[min(v['yield'] for v in params.values()),max(v['yield'] for v in params.values())],
        'cost_range':[min(v['cost'] for v in params.values()),max(v['cost'] for v in params.values())],
        'demand_field_in_headers':any('销售量' in str(param_ws.cell(1,c).value) for c in range(1,9))}

    # Enumerate eligible keys ONLY from Attachment 1 explicit notes for coverage diagnosis.
    grain=[k for k,v in crops.items() if v['type'].startswith('粮食') and v['name']!='水稻']
    veg=[k for k,v in crops.items() if v['type'].startswith('蔬菜') and v['name'] not in ('大白菜','白萝卜','红萝卜')]
    winter=[k for k,v in crops.items() if v['name'] in ('大白菜','白萝卜','红萝卜')]
    fungi=[k for k,v in crops.items() if v['type']=='食用菌']
    rice=[k for k,v in crops.items() if v['name']=='水稻']
    eligible=[]
    for lt in ('平旱地','梯田','山坡地'):
        eligible.extend((k,lt,'单季') for k in grain)
    eligible.extend((k,'水浇地','单季') for k in rice)
    for lt in ('水浇地','普通大棚','智慧大棚'):
        eligible.extend((k,lt,'第一季') for k in veg)
    eligible.extend((k,'水浇地','第二季') for k in winter)
    eligible.extend((k,'普通大棚','第二季') for k in fungi)
    eligible.extend((k,'智慧大棚','第二季') for k in veg)
    unresolved=[];inherited=[]
    for k in eligible:
        if k in params:continue
        if k[1]=='智慧大棚' and k[2]=='第一季' and (k[0],'普通大棚','第一季') in params:inherited.append(k)
        else:unresolved.append(k)
    out['eligible_parameter_coverage_diagnostic']={'eligible_type_crop_season_count':len(eligible),'direct_count':sum(k in params for k in eligible),
        'explicit_note_count':len(inherited),'unresolved':unresolved,'parameter_keys_outside_explicit_rules':[k for k in params if k not in eligible],
        'planting_records_outside_explicit_rules':[p for p in planting if (p['crop_id'],lands[p['land']]['type'],p['season']) not in eligible]}

    eligible_pairs={(k[0],k[2]) for k in eligible}
    observed_pairs={(p['crop_id'],p['season']) for p in planting}
    season_prices=defaultdict(set)
    for (cid,lt,season),v in params.items():season_prices[(cid,season)].add(v['price'])
    out['demand_identifiability']={'eligible_crop_season_pairs':len(eligible_pairs),'observed_crop_season_pairs':len(observed_pairs),
        'unobserved_eligible_pairs':[{'crop_id':cid,'crop':crops[cid]['name'],'season':season} for cid,season in sorted(eligible_pairs-observed_pairs)],
        'interpretation':'No historical planting is not evidence of zero market demand. No demand values are constructed in this audit.',
        'same_crop_season_price_disagreements':{str(k):sorted(v) for k,v in season_prices.items() if len(v)>1}}

    # Independent numerical check from raw OOXML, without openpyxl or merged lookup.
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    def xml_numeric(path, sheet, column):
        with zipfile.ZipFile(path) as z:
            root=ET.fromstring(z.read(f'xl/worksheets/sheet{sheet}.xml'))
        result=[]
        for c in root.findall('.//m:sheetData/m:row/m:c',ns):
            if re.fullmatch(column+r'\d+',c.attrib['r']) and c.attrib.get('t','n')=='n':
                v=c.find('m:v',ns)
                if v is not None and v.text is not None:result.append(Decimal(v.text))
        return result
    raw_land=xml_numeric(SOURCE/'附件1.xlsx',1,'C')
    raw_crop=xml_numeric(SOURCE/'附件1.xlsx',2,'A')
    raw_plant=xml_numeric(SOURCE/'附件2.xlsx',1,'E')
    raw_param=xml_numeric(SOURCE/'附件2.xlsx',2,'A')
    independent={'land_count':len(raw_land),'physical_area':str(sum(raw_land)),
        'crop_count':len(raw_crop),'planting_count':len(raw_plant),'sown_season_area':str(sum(raw_plant)),
        'parameter_count':len(raw_param)}
    independent['agrees_with_openpyxl']=(len(raw_land)==len(land_rows) and len(raw_crop)==len(crop_rows)
        and len(raw_plant)==len(pr) and len(raw_param)==len(rr)
        and abs(float(sum(raw_land))-sum(v['area'] for v in lands.values()))<1e-8
        and abs(float(sum(raw_plant))-sum(p['area'] for p in planting))<1e-8)
    assert independent['agrees_with_openpyxl']
    out['independent_ooxml_verification']=independent

    templates={}
    for name in ('result1_1.xlsx','result1_2.xlsx','result2.xlsx'):
        w=books[name];sigs=[];yearmeta={}
        for ws in w:
            labels={'headers':[ws.cell(1,c).value for c in range(2,44)],'lands':[ws.cell(r,2).value for r in range(2,84)],
                    'season_labels':[ws.cell(r,1).value for r in (2,56)],'notes':[ws.cell(r,2).value for r in (85,86,87)]}
            sigs.append(json.dumps(labels,ensure_ascii=False,sort_keys=True))
            vals=[c for row in ws.iter_rows(min_row=2,max_row=83,min_col=3,max_col=43) for c in row if c.value is not None]
            yearmeta[ws.title]={'nonempty_result_cells':len(vals),'examples':[{'cell':c.coordinate,'value':c.value} for c in vals[:6]],
                'header_crop_names_match':labels['headers'][1:]==[crop_ws.cell(r,2).value for r in crop_rows],
                'validation_count':len(ws.data_validations.dataValidation),
                'result_style_ids':dict(Counter(c.style_id for row in ws.iter_rows(min_row=2,max_row=83,min_col=3,max_col=43) for c in row))}
        templates[name]={'years':w.sheetnames,'labels_consistent_across_years':len(set(sigs))==1,'by_year':yearmeta}
    out['templates']=templates
    for w in books.values():w.close()
    out['source_hashes_after']={str(p.relative_to(SOURCE)):sha(p) for p in paths}
    out['source_unchanged']=out['source_hashes_before']==out['source_hashes_after']
    assert out['source_unchanged']
    OUT.mkdir(parents=True,exist_ok=True)
    (OUT/'audit.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
    compact={k:v for k,v in out.items() if k not in ('books','templates','source_hashes_before','source_hashes_after')}
    compact['templates']={k:{'years':v['years'],'labels_consistent_across_years':v['labels_consistent_across_years'],
        'nonempty_result_cells':sum(z['nonempty_result_cells'] for z in v['by_year'].values()),
        'header_matches_all_years':all(z['header_crop_names_match'] for z in v['by_year'].values())} for k,v in templates.items()}
    print(json.dumps(compact,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
