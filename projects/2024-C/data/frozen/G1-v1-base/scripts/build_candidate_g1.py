"""U003 / G1-v1 approved data preparation; no demand construction or modeling."""
from pathlib import Path
from decimal import Decimal
from collections import defaultdict
from datetime import datetime, timezone
import csv
import hashlib
import json
import re
import warnings
import openpyxl

P=Path(__file__).resolve().parents[1]
SRC=P.parents[1]/'培训作业'/'第四次培训作业'
OUT=P/'data/candidates/G1-v1'
LOG=[]
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def js(x):return json.dumps(x,ensure_ascii=False,separators=(',',':'))
def dec(x):return Decimal(str(x))
def num(x):return format(dec(x).normalize(),'f')
def event(kind,record,**kw):LOG.append({'operation':kind,'record_id':record,**kw})

def source_row(ws,r,fields,file,record):
    values={}; provenance={}
    for col,name in enumerate(fields,1):
        cell=ws.cell(r,col); anchor=cell; merged=None
        for m in ws.merged_cells.ranges:
            if m.min_row<=r<=m.max_row and m.min_col<=col<=m.max_col:
                merged=str(m);anchor=ws.cell(m.min_row,m.min_col);break
        val=anchor.value
        provenance[name]={'cell':cell.coordinate,'raw_value':cell.value,'anchor':anchor.coordinate,
                          'anchor_value':val,'merged_range':merged}
        if cell.coordinate!=anchor.coordinate:
            event('P02_merged_anchor',record,field=name,source_file=file,sheet=ws.title,
                  cell=cell.coordinate,anchor=anchor.coordinate,merged_range=merged,raw_value=cell.value,value=val)
        # Only identifiers/names/categories/season labels are trimmed; source prose is preserved.
        if name in {'land_id','land_type','crop_name','crop_type','season'} and isinstance(val,str):
            clean=val.strip()
            if clean!=val:event('P02_trim_edges',record,field=name,source_file=file,sheet=ws.title,cell=anchor.coordinate,raw_value=val,value=clean)
            val=clean
        values[name]=val
    values.update(record_id=record,source_file=file,source_sheet=ws.title,source_row=r,source_cells_json=js(provenance))
    return values

def write_csv(name,rows):
    assert rows
    with (OUT/name).open('x',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)

def main():
    state=json.loads((P/'state/state.json').read_text(encoding='utf-8'))
    assert state['stage']=='DATA_EXEC' and any(a['id']=='U003' and a['plan_version']=='G1-v1' for a in state['authorizations'])
    assert not OUT.exists(),'Candidate directory exists; inspect before any recovery, never overwrite.'
    audit=json.loads((P/'audit/G1-v1/audit.json').read_text(encoding='utf-8'))
    hashes=audit['source_hashes_before']
    assert all(sha(SRC/f)==h for f,h in hashes.items()),'Source version changed'
    with warnings.catch_warnings(record=True) as warns:
        warnings.simplefilter('always')
        w1=openpyxl.load_workbook(SRC/'附件1.xlsx',data_only=False)
        w2=openpyxl.load_workbook(SRC/'附件2.xlsx',data_only=False)
    lws,cws=w1.worksheets;hws,pws=w2.worksheets
    lands=[];crops=[];history=[];original=[]
    for r in range(2,lws.max_row+1):
        if re.fullmatch('[A-F][0-9]+',str(lws.cell(r,1).value)):
            d=source_row(lws,r,['land_id','land_type','area_mu'],'附件1.xlsx',f'L{r:03}')
            d['area_mu']=num(d['area_mu']);lands.append(d)
    for r in range(2,cws.max_row+1):
        if isinstance(cws.cell(r,1).value,int):
            d=source_row(cws,r,['crop_id','crop_name','crop_type','suitable_land_text','notes'],'附件1.xlsx',f'C{r:03}')
            d['is_legume']='true' if '豆类' in d['crop_type'] else 'false';crops.append(d)
    for r in range(2,hws.max_row+1):
        if isinstance(hws.cell(r,2).value,int):
            d=source_row(hws,r,['land_id','crop_id','crop_name','crop_type','area_mu','season'],'附件2.xlsx',f'H{r:03}')
            d['area_mu']=num(d['area_mu']);d['planting_year']=2023;d['template_season']='第一季' if d['season']=='单季' else d['season'];history.append(d)
    for r in range(2,pws.max_row+1):
        if isinstance(pws.cell(r,1).value,int):
            d=source_row(pws,r,['source_sequence','crop_id','crop_name','land_type','season','yield_jin_per_mu','cost_yuan_per_mu','price_interval_raw'],'附件2.xlsx',f'P{r:03}')
            m=re.fullmatch(r'(\d+(?:\.\d+)?)\s*-\s*(\d+(?:\.\d+)?)',d['price_interval_raw']);assert m
            lo,hi=dec(m[1]),dec(m[2]);assert 0<lo<=hi
            d['yield_jin_per_mu']=num(d['yield_jin_per_mu']);d['cost_yuan_per_mu']=num(d['cost_yuan_per_mu'])
            d.update(price_lower_yuan_per_jin=num(lo),price_upper_yuan_per_jin=num(hi),price_midpoint_candidate_yuan_per_jin=num((lo+hi)/2),price_status='interval_observed_midpoint_assumed_candidate')
            event('P04_interval_and_midpoint',d['record_id'],source=f'附件2.xlsx/{pws.title}!H{r}',raw=d['price_interval_raw'],lower=num(lo),upper=num(hi),midpoint_candidate=num((lo+hi)/2));original.append(d)
    effective=[]
    for d in original:
        x=dict(d);x.update(origin='source',derived_from_record_id='',rule_source='');effective.append(x)
        if d['land_type']=='普通大棚' and d['season']=='第一季':
            x=dict(d);x.update(record_id=f"I{d['crop_id']:02}",land_type='智慧大棚',origin='derived_from_note',derived_from_record_id=d['record_id'],rule_source=f'附件2.xlsx/{pws.title}!B111')
            effective.append(x);event('P03_explicit_note_reference',x['record_id'],from_record=d['record_id'],from_land_type='普通大棚',to_land_type='智慧大棚',rule_source=x['rule_source'])
    ld={d['land_id']:d for d in lands};cd={d['crop_id']:d for d in crops}
    pd={(d['crop_id'],d['land_type'],d['season']):d for d in effective}
    # Independently list admissible type/crop/season keys from Attachment 1 notes.
    allowed=set()
    for c in crops:
        i=c['crop_id'];t=c['crop_type'];name=c['crop_name']
        if t.startswith('粮食'):
            allowed.update((i,lt,'单季') for lt in (['水浇地'] if name=='水稻' else ['平旱地','梯田','山坡地']))
        elif t.startswith('蔬菜'):
            if name in ('大白菜','白萝卜','红萝卜'):allowed.add((i,'水浇地','第二季'))
            else:
                allowed.update((i,lt,'第一季') for lt in ('水浇地','普通大棚','智慧大棚'));allowed.add((i,'智慧大棚','第二季'))
        elif t=='食用菌':allowed.add((i,'普通大棚','第二季'))
        else:raise ValueError(t)
    assert allowed==set(pd) and len(pd)==125
    aggregates=defaultdict(list)
    for h in history:
        key=(h['crop_id'],ld[h['land_id']]['land_type'],h['season']);p=pd[key]
        prod=dec(h['area_mu'])*dec(p['yield_jin_per_mu'])
        h.update(parameter_record_id=p['record_id'],production_estimate_jin=num(prod),production_status='area_times_statistical_yield_not_observed_sales')
        aggregates[(h['crop_id'],h['season'])].append(h)
        event('P05_production_estimate',h['record_id'],area_mu=h['area_mu'],parameter_record_id=p['record_id'],yield_jin_per_mu=p['yield_jin_per_mu'],production_estimate_jin=num(prod))
    production=[]
    for (cid,season),hs in sorted(aggregates.items()):
        production.append({'crop_id':cid,'crop_name':cd[cid]['crop_name'],'planting_year':2023,'season':season,
            'production_estimate_jin':num(sum((dec(h['production_estimate_jin']) for h in hs),Decimal(0))),
            'planted_area_mu':num(sum((dec(h['area_mu']) for h in hs),Decimal(0))),
            'contributions_json':js([{'history_record_id':h['record_id'],'parameter_record_id':h['parameter_record_id'],'area_mu':h['area_mu'],'production_estimate_jin':h['production_estimate_jin']} for h in hs]),
            'interpretation':'historical_production_estimate_not_sales_or_demand'})
    sums={(p['crop_id'],p['season']):p['production_estimate_jin'] for p in production}
    coverage=[]
    for cid,season in sorted({(k[0],k[2]) for k in allowed}):
        coverage.append({'crop_id':cid,'crop_name':cd[cid]['crop_name'],'season':season,
            'historical_planting_status':'observed' if (cid,season) in sums else 'unobserved',
            'production_estimate_jin':sums.get((cid,season),''),'normal_sales_limit_jin':'',
            'demand_status':'unknown_not_imputed','source_history_ids_json':js([h['record_id'] for h in aggregates.get((cid,season),[])]),
            'eligible_parameter_ids_json':js(sorted(p['record_id'] for k,p in pd.items() if k[0]==cid and k[2]==season))})
    mapping={'scope':'Q1 templates only; no workbook writing','years':list(range(2024,2031)),
        'source_season_to_template_season':{'单季':'第一季','第一季':'第一季','第二季':'第二季'},
        'calendar_policy':'Preserve source planting-year/season labels; do not infer harvest-year demand pools.',
        'rules_source':'附件1.xlsx/乡村种植的农作物!E2,B44:B46; 附件2.xlsx/2023年统计的相关数据!B111',
        'allowed_parameter_keys':[{'crop_id':i,'land_type':lt,'season':s} for i,lt,s in sorted(allowed)],'templates':{}}
    for fn in ('result1_1.xlsx','result1_2.xlsx'):
        with warnings.catch_warnings():
            warnings.simplefilter('ignore');tw=openpyxl.load_workbook(SRC/'附件3'/fn)
        annual={}
        for ws in tw:
            cols={str(cid):openpyxl.utils.get_column_letter(next(c for c in range(3,44) if ws.cell(1,c).value.strip()==cd[cid]['crop_name'])) for cid in cd}
            rows={}
            for r in range(2,84):
                season='第一季' if r<56 else '第二季';rows.setdefault(season,{})[ws.cell(r,2).value]=r
            annual[ws.title]={'crop_columns':cols,'land_rows_by_template_season':rows,'result_range':'C2:AQ83'}
        mapping['templates'][fn]=annual;tw.close()
    assert all(sha(SRC/f)==h for f,h in hashes.items())
    OUT.mkdir(parents=True)
    tables={'lands.csv':lands,'crops.csv':crops,'planting_2023.csv':history,'parameters_original.csv':original,'parameters_effective.csv':effective,'production_2023.csv':production,'coverage.csv':coverage}
    for fn,data in tables.items():write_csv(fn,data)
    (OUT/'template_mapping.json').write_text(json.dumps(mapping,ensure_ascii=False,indent=2),encoding='utf-8')
    (OUT/'transform_log.jsonl').write_text('\n'.join(js(x) for x in LOG)+'\n',encoding='utf-8')
    manifest={'candidate_id':'G1-v1','status':'CANDIDATE_NOT_FROZEN','authorization_id':'U003',
        'created_at_utc':datetime.now(timezone.utc).isoformat(),'source_directory':str(SRC),'source_hashes':hashes,
        'builder':{'path':'scripts/build_candidate_g1.py','sha256':sha(Path(__file__))},'openpyxl_version':openpyxl.__version__,
        'units':{'area_mu':'亩','yield_jin_per_mu':'斤/亩','cost_yuan_per_mu':'元/亩','prices':'元/斤','production_estimate_jin':'斤'},
        'row_counts':{fn:len(data) for fn,data in tables.items()},
        'files':{p.name:sha(p) for p in sorted(OUT.iterdir())},
        'transform_event_count':len(LOG),'load_warnings':sorted({str(w.message) for w in warns}),
        'unknown_policy':'All 59 normal sales limits unknown. 12 unobserved crop-season production estimates blank. No blank-to-zero conversion.',
        'boundaries':['No demand proxy adopted','No spatial bean-rotation assumptions','No simulations or optimization','No source/template writes','No human verification claimed']}
    (OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(js({'candidate':str(OUT),'row_counts':manifest['row_counts'],'transform_event_count':len(LOG),'source_unchanged':True}))
    w1.close();w2.close()

if __name__=='__main__':main()
