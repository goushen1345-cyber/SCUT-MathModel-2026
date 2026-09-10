"""AUDIT-REVISIT-v1: source diagnostics only, no processed dataset or model outputs.
No workbook save, no production/revenue/demand calculation, no EDA or research imports.
"""
from pathlib import Path
from collections import Counter, defaultdict
from decimal import Decimal
from datetime import datetime, timezone
import csv, hashlib, json, re, sys, warnings
from openpyxl import load_workbook
from pypdf import PdfReader

HERE=Path(__file__).resolve().parent
PROJECT=HERE.parents[1]
ROOT=PROJECT.parents[1]
RAW=ROOT/'培训作业/第四次培训作业'
FROZEN=PROJECT/'data/frozen/G1-v1-base'
START='2026-09-10T07:29:49+00:00'
def dump(name,obj):
    (HERE/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf-8')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def clean(v):return v.strip() if isinstance(v,str) else v
def dec(v):return Decimal(str(v))
def anchor(s,c):
    for m in s.merged_cells.ranges:
        if c.coordinate in m:
            a=s.cell(m.min_row,m.min_col)
            return a.value,a.coordinate,str(m)
    return c.value,c.coordinate,None
def ref(s,r,cols):return s.title+'!'+cols[0]+str(r)+':'+cols[-1]+str(r)
def duplicates(items):return [{'key':str(k),'count':n} for k,n in Counter(items).items() if n>1]
def csvrows(n):
    with (FROZEN/n).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))

def main():
    if (HERE/'diagnostics.json').exists():raise FileExistsError('Use new diagnostic output version, do not silently overwrite')
    if (datetime.now(timezone.utc)-datetime.fromisoformat(START)).total_seconds()>1800:raise TimeoutError('Approved 30 minute budget expired')
    state=json.loads((PROJECT/'state/state.json').read_text(encoding='utf-8'))
    assert state['stage']=='DATA_PLAN' and state['authorizations'][-1]['id']=='U029'
    protected=set(RAW.rglob('*.xlsx'))|{RAW/'C题.pdf'}
    def snapshots(x):
        if isinstance(x,dict):
            for k,v in x.items():
                if isinstance(v,str) and re.fullmatch('[0-9a-f]{64}',v) and (PROJECT/k).is_file():protected.add(PROJECT/k)
                else:snapshots(v)
        elif isinstance(x,list):
            for v in x:snapshots(v)
    snapshots(state)
    for sub in ['paper/final','paper/final_v2','paper/understanding_v2','paper/figures/model_structure_v3']:
        protected.update(p for p in (PROJECT/sub).rglob('*') if p.is_file())
    protected.add(PROJECT/'paper/front_matter_understanding_v2.md')
    before={str(p.relative_to(ROOT)).replace('\\','/'):sha(p) for p in sorted(protected)}
    dump('protected_before.json',before)
    inventory=[]; books={}; warning_log=[]
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter('always')
        for p in sorted(RAW.rglob('*.xlsx')):books[str(p.relative_to(RAW))]=load_workbook(p,data_only=False)
        warning_log=[str(w.message) for w in caught]
    for filename,w in books.items():
        for s in w:
            inventory.append({'file':filename,'sheet':s.title,'dimension':s.calculate_dimension(),
              'headers':[c.value for c in s[1]],'merged_ranges':[str(x) for x in s.merged_cells.ranges],
              'nonempty_cells':sum(c.value is not None for row in s for c in row),
              'formula_cells':[c.coordinate for row in s for c in row if c.data_type=='f']})
    problem=PdfReader(RAW/'C题.pdf')
    (HERE/'problem_source.txt').write_text('\n'.join(p.extract_text() for p in problem.pages),encoding='utf-8')
    l=books['附件1.xlsx']['乡村的现有耕地'];c=books['附件1.xlsx']['乡村种植的农作物']
    h=books['附件2.xlsx']['2023年的农作物种植情况'];p=books['附件2.xlsx']['2023年统计的相关数据']
    land_rows=[r for r in range(2,l.max_row+1) if l.cell(r,1).value is not None]
    crop_rows=[r for r in range(2,c.max_row+1) if isinstance(c.cell(r,1).value,int)]
    hist_rows=[r for r in range(2,h.max_row+1) if isinstance(h.cell(r,2).value,int)]
    par_rows=[r for r in range(2,p.max_row+1) if isinstance(p.cell(r,1).value,int)]
    lands={clean(l.cell(r,1).value):(clean(l.cell(r,2).value),dec(l.cell(r,3).value),r) for r in land_rows}
    crops={c.cell(r,1).value:(clean(c.cell(r,2).value),clean(c.cell(r,3).value),r) for r in crop_rows}
    whitespace=[]; blanks=[]; malformed=[]
    for s,rows,cols in [(l,land_rows,range(1,4)),(c,crop_rows,range(1,5)),(h,hist_rows,range(1,7)),(p,par_rows,range(1,9))]:
        for r in rows:
            for col in cols:
                cell=s.cell(r,col);v=cell.value;a,coord,m=anchor(s,cell)
                if isinstance(v,str) and v!=v.strip():whitespace.append({'source':s.title+'!'+cell.coordinate,'raw':v,'stripped_for_comparison_only':v.strip()})
                if v is None:blanks.append({'source':s.title+'!'+cell.coordinate,'anchor':coord,'merged_range':m,'anchor_nonempty':a is not None})
    def allowed(cid):
        name,typ,_=crops[cid]
        if name=='水稻':return {('水浇地','单季')}
        if typ.startswith('粮食'):return {(x,'单季') for x in ['平旱地','梯田','山坡地']}
        if name in ['大白菜','白萝卜','红萝卜']:return {('水浇地','第二季')}
        if typ.startswith('蔬菜'):return {('水浇地','第一季'),('普通大棚','第一季'),('智慧大棚','第一季'),('智慧大棚','第二季')}
        if typ=='食用菌':return {('普通大棚','第二季')}
        raise ValueError('Unrecognized source crop type')
    par_index=defaultdict(list);bad_names=[];price_checks=[]
    for r in par_rows:
        cid=p.cell(r,2).value;name=clean(p.cell(r,3).value);lt=clean(p.cell(r,4).value);ss=clean(p.cell(r,5).value)
        par_index[(cid,lt,ss)].append(r)
        if cid not in crops or name!=crops[cid][0]:bad_names.append(ref(p,r,'BC'))
        for col in [6,7]:
            v=p.cell(r,col).value
            if not isinstance(v,(int,float)) or dec(v)<=0:malformed.append({'source':p.cell(r,col).coordinate,'value':v,'reason':'nonpositive_or_not_numeric'})
        raw=p.cell(r,8).value;m=re.fullmatch(r'\s*(\d+(?:\.\d+)?)\s*[-–—]\s*(\d+(?:\.\d+)?)\s*',str(raw))
        valid=bool(m) and dec(m[1])>0 and dec(m[2])>=dec(m[1])
        price_checks.append({'source':'2023年统计的相关数据!H'+str(r),'raw':raw,'valid_positive_ordered_interval':valid})
    expected={(cid,lt,ss) for cid in crops for lt,ss in allowed(cid)}
    missing=sorted(expected-set(par_index));extra=sorted(set(par_index)-expected)
    notes={k:s[k].value for s,keys in [(l,['D2']),(c,['E2','B44','B45','B46']),(p,['B110','B111'])] for k in []}
    notes={s.title+'!'+a:s[a].value for s,addresses in [(l,['D2']),(c,['E2','B44','B45','B46']),(p,['B110','B111'])] for a in addresses}
    hist_checks=[];areas=defaultdict(Decimal);source_rows=defaultdict(list);histkeys=[];observed=set()
    for r in hist_rows:
        land,anc,m=anchor(h,h.cell(r,1));land=clean(land);cid=h.cell(r,2).value;ss=clean(h.cell(r,6).value)
        name=clean(h.cell(r,3).value);typ=clean(h.cell(r,4).value);area=h.cell(r,5).value
        lt=lands[land][0] if land in lands else None
        key=(cid,lt,ss);rows=par_index.get(key,[]);rule_rows=[]
        if not rows and lt=='智慧大棚' and ss=='第一季':rule_rows=par_index.get((cid,'普通大棚','第一季'),[])
        histkeys.append((land,cid,ss));observed.add((cid,'第一季' if ss=='单季' else ss))
        chk={'source':ref(h,r,'AF'),'land_anchor':anc,'merge':m,'land_known':land in lands,'crop_known':cid in crops,
             'name_type_match':cid in crops and (name,typ)==crops[cid][:2],'suitable':cid in crops and (lt,ss) in allowed(cid),
             'direct_parameter_rows':rows,'rule_parameter_rows':rule_rows,'rule_source':'2023年统计的相关数据!B111' if rule_rows else None}
        hist_checks.append(chk)
        if not isinstance(area,(int,float)) or dec(area)<=0:malformed.append({'source':ref(h,r,'EE'),'value':area,'reason':'area_nonpositive_or_not_numeric'})
        else:areas[(land,ss)]+=dec(area);source_rows[(land,ss)].append(r)
    area_checks=[{'land':land,'season':ss,'source_rows':source_rows[(land,ss)],'sum_original_area_mu':str(a),'capacity_mu':str(lands[land][1]),
       'difference_mu':str(a-lands[land][1]),'over_capacity':a>lands[land][1]} for (land,ss),a in areas.items()]
    legal_seasons={(cid,'第一季' if ss=='单季' else ss) for cid,lt,ss in expected}
    unobserved=sorted(legal_seasons-observed)
    extrema=[]
    for col in [6,7]:
        valid=[r for r in par_rows if isinstance(p.cell(r,col).value,(int,float))]
        for tag,val in [('min',min(p.cell(r,col).value for r in valid)),('max',max(p.cell(r,col).value for r in valid))]:
            extrema.append({'field':p.cell(1,col).value,'kind':tag,'value':val,'rows':[{'source':ref(p,r,'BH'),'crop':p.cell(r,3).value,'land_type':p.cell(r,4).value,'season':p.cell(r,5).value} for r in valid if p.cell(r,col).value==val], 'status':'raw_extreme_for_context_review_not_error_label'})
    template_checks=[]
    for filename,w in books.items():
        if not filename.startswith('附件3'):continue
        base=None
        for s in w:
            headings=[clean(s.cell(1,col).value) for col in range(3,44)]
            labels=[];badland=[]
            for rr in range(2,84):
                land=clean(s.cell(rr,2).value)
                if land is not None:
                    sea,anc,m=anchor(s,s.cell(rr,1));sea=re.sub(r'\s','',str(sea))
                    labels.append((rr,land,sea))
                    if land not in lands:badland.append({'row':rr,'land':land})
            structural={'headers':headings,'row_labels':labels,'merges':sorted(str(x) for x in s.merged_cells.ranges),'notes':[s[f'B{x}'].value for x in [85,86,87]]}
            if base is None:base=structural
            populated=[cell.coordinate for row in s.iter_rows(min_row=2,max_row=83,min_col=3,max_col=43) for cell in row if cell.value is not None]
            template_checks.append({'file':filename,'sheet':s.title,'year_set_correct':w.sheetnames==[str(t) for t in range(2024,2031)],
              'same_structure_as_first_year':structural==base,'crop_headers_match':[crops[k][0] for k in sorted(crops)]==headings,
              'row_count':len(labels),'unknown_lands':badland,'season_counts':dict(Counter(x[2] for x in labels)),
              'duplicate_land_season':duplicates((x[1],x[2]) for x in labels),'nonempty_result_cells':populated,'notes':structural['notes']})
    # Compare only source-backed fields, never calculate/use old production or midpoint columns.
    old_checks=[]
    specs=[('lands.csv',l,{'land_id':1,'land_type':2,'area_mu':3}),('crops.csv',c,{'crop_id':1,'crop_name':2,'crop_type':3,'suitable_land_text':4}),
      ('planting_2023.csv',h,{'land_id':1,'crop_id':2,'crop_name':3,'crop_type':4,'area_mu':5,'season':6}),
      ('parameters_original.csv',p,{'source_sequence':1,'crop_id':2,'crop_name':3,'land_type':4,'season':5,'yield_jin_per_mu':6,'cost_yuan_per_mu':7,'price_interval_raw':8})]
    for filename,s,mapping in specs:
        records=csvrows(filename);mismatches=[];checked=0
        for row in records:
            rr=int(row['source_row'])
            for field,col in mapping.items():
                v,anc,mer=anchor(s,s.cell(rr,col));v=clean(v);old=row[field];checked+=1
                match=dec(old)==dec(v) if isinstance(v,(int,float)) else clean(old)==v
                if not match:mismatches.append({'source':s.title+'!'+s.cell(rr,col).coordinate,'field':field,'source_value':v,'old_value':old})
        old_checks.append({'file':filename,'records':len(records),'source_fields_compared':checked,'mismatches':mismatches})
    effective=csvrows('parameters_effective.csv');rule_check=[]
    for row in effective:
        if row['origin']=='source':continue
        cid=int(row['crop_id']);sr=par_index.get((cid,'普通大棚','第一季'),[])
        match=len(sr)==1 and all(dec(row[f])==dec(p.cell(sr[0],col).value) for f,col in [('yield_jin_per_mu',6),('cost_yuan_per_mu',7)]) and row['price_interval_raw']==p.cell(sr[0],8).value
        rule_check.append({'record_id':row['record_id'],'crop_id':cid,'target_land':row['land_type'],'season':row['season'],'ordinary_source_rows':sr,'raw_parameter_values_match':match,'declared_rule_source':row['rule_source']})
    summary={'land_records':len(land_rows),'crop_records':len(crop_rows),'history_records':len(hist_rows),'original_parameter_records':len(par_rows),
      'blank_cells_in_selected_fields':len(blanks),'merged_blanks_with_nonempty_anchor':sum(x['merged_range'] is not None and x['anchor_nonempty'] for x in blanks),
      'unexplained_blank_cells':[x for x in blanks if not x['merged_range'] or not x['anchor_nonempty']], 'whitespace_cells':len(whitespace),
      'duplicate_land_keys':duplicates(clean(l.cell(r,1).value) for r in land_rows),'duplicate_crop_keys':duplicates(c.cell(r,1).value for r in crop_rows),
      'duplicate_history_keys':duplicates(histkeys),'duplicate_parameter_keys':duplicates(k for k,rs in par_index.items() for _ in rs),
      'parameter_crop_name_mismatches':bad_names,'parameter_invalid_suitability':extra,'missing_direct_parameter_keys':missing,
      'history_direct_parameter_matches':sum(len(x['direct_parameter_rows'])==1 for x in hist_checks),'history_rule_parameter_matches':sum(len(x['rule_parameter_rows'])==1 for x in hist_checks),
      'history_failed_checks':[x for x in hist_checks if not all(x[k] for k in ['land_known','crop_known','name_type_match','suitable']) or len(x['direct_parameter_rows'])+len(x['rule_parameter_rows'])!=1],
      'area_groups_checked':len(area_checks),'area_over_capacity':[x for x in area_checks if x['over_capacity']],'area_not_equal_capacity':[x for x in area_checks if dec(x['difference_mu'])!=0],
      'price_interval_count':len(price_checks),'invalid_price_intervals':[x for x in price_checks if not x['valid_positive_ordered_interval']], 'invalid_numeric_cells':malformed,
      'legal_crop_output_season_pairs':len(legal_seasons),'observed_crop_output_season_pairs':len(observed),'unobserved_legal_pairs':[{'crop_id':cid,'crop':crops[cid][0],'template_season':ss,'meaning':'not planted in 2023 records, not zero demand'} for cid,ss in unobserved],
      'template_sheets_checked':len(template_checks),'frozen_source_comparisons':old_checks,'frozen_rule_parameter_checks':rule_check}
    for w in books.values():w.close()
    changed=[f for f,digest in before.items() if sha(ROOT/f)!=digest]
    dump('inventory.json',{'pdf_pages':len(problem.pages),'sheets':inventory,'read_warnings':warning_log})
    dump('diagnostic_details.json',{'whitespace':whitespace,'blank_sources':blanks,'price_interval_checks':price_checks,'history_key_checks':hist_checks,'area_capacity_checks':area_checks,'raw_extreme_review':extrema,'template_checks':template_checks,'source_notes':notes})
    dump('diagnostics.json',summary)
    dump('execution.json',{'version':'AUDIT-REVISIT-v1','authorization':'U029','started_utc':START,'diagnostics_finished_utc':datetime.now(timezone.utc).isoformat(),
      'python':sys.version,'source_code_sha256':sha(Path(__file__)),'protected_files':len(before),'changed_protected_files':changed,
      'scope':'diagnostic only; no input export, cleaning, EDA, production/demand/revenue/midpoint computation, model or experiment',
      'budget_seconds':1800,'command':'bundled Python audit/AUDIT-REVISIT-v1/audit_readonly.py'})
    print(json.dumps(summary,ensure_ascii=False,default=str))

if __name__=='__main__':main()
