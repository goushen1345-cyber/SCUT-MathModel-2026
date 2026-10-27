from pathlib import Path
import ast,shutil,datetime,json
import openpyxl
from openpyxl.styles import Alignment, Font, Border, Side, PatternFill

P=Path(__file__).resolve().parents[2]
D=P/'支撑材料';C=P/'paper/提交配套'
T=Path(__file__).parent/'presentation_update'
T.mkdir(exist_ok=True)
B=T/'before';B.mkdir(exist_ok=True)

# Load only the style function; no numerical routines or exports are run.
source=(D/'code/export_results.py').read_text('utf-8')
tree=ast.parse(source)
fn=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=='format_workbook')
exec(compile(ast.Module(body=[fn],type_ignores=[]),'format_workbook','exec'))

def snapshot(w):
 return [(s.title,s.max_row,s.max_column,str(s.merged_cells),[(c.coordinate,c.value,c.data_type,c.number_format) for row in s for c in row]) for s in w]

checks=[]
for path in sorted(D.glob('result*.xlsx')):
 backup=B/path.name
 if not backup.exists():shutil.copy2(path,backup)
 w=openpyxl.load_workbook(path)
 before=snapshot(w)
 format_workbook(w)
 output=T/path.name
 w.save(output)
 verified=openpyxl.load_workbook(output)
 assert snapshot(verified)==before,path.name
 assert all(c.font.name=='宋体' and c.font.sz==11 for s in verified for row in s for c in row)
 try:
  shutil.copy2(output,path)
  status='updated'
 except PermissionError:status='locked; styled copy in presentation_update'
 checks.append(dict(file=path.name,status=status,values_and_number_formats_unchanged=True))

for name in ['AI工具使用详情.tex','AI工具使用详情.pdf']:
 if not (B/name).exists():shutil.copy2(C/name,B/name)
s=(C/'AI工具使用详情.tex').read_text('utf-8')
needle='OpenAI Codex & GPT-6（按本项目会话标识；精确部署子版本未留存）'
assert needle in s
s=s.replace(needle, 'DeepSeek & 具体版本未提供 & 队伍确认：已使用DeepSeek另行完成本项目中由Codex辅助的全部工作，涵盖题意、数据、建模、代码、验证、图表及论文等环节。\\\\\nOpenAI Codex & GPT-6（按本项目会话标识；精确部署子版本未留存）')
s=s.replace('本报告依据2026年9月10日至13日本项目工作聊天及保存的使用记录编写，记录对象为', '本报告依据2026年9月10日至13日本项目工作聊天、保存的使用记录及队伍对DeepSeek使用情况的补充确认编写，记录对象为')
s=s.replace('以下交互根据本项目真实聊天节选或摘要整理，回答为内容概述，并非逐字完整对话。','以下交互根据本项目可追溯的Codex聊天节选或摘要整理，回答为内容概述，并非逐字完整对话。队伍确认已用DeepSeek完成相同范围的工作；其原始对话未纳入当前项目记录，故不将下列交互改署为DeepSeek对话。')
start=s.index('\\textbf{报告覆盖范围。}')
end=s.index('\\end{document}',start)
s=s[:start]+'\\textbf{报告覆盖范围。}队伍已确认，Codex辅助完成的工作均另行使用DeepSeek做过一遍，本报告据此补充DeepSeek的全流程使用情况。DeepSeek的具体版本、逐次使用日期和原始对话未在当前项目记录中提供，故不以官网当前版本代替实际使用版本。Codex部分依据本项目留存记录汇总。组员其他聊天材料及外部立体示意图的制作工具与版本仍以其真实来源为准。\n\n'+s[end:]
(C/'AI工具使用详情.tex').write_text(s,encoding='utf-8')
p=C/'提交前待补事项.md'
s=p.read_text('utf-8')
lines=s.splitlines()
lines=[('1. 队伍已确认：已使用DeepSeek重新完成Codex辅助的全部工作。AI报告已补入DeepSeek及全流程使用范围，并保留Codex真实记录；DeepSeek具体版本和逐次日期尚未提供，未臆填最新版。提交前由队伍核对报告并补充可查的原始对话。' if l.startswith('1. ') else l) for l in lines]
p.write_text('\n'.join(lines)+'\n',encoding='utf-8')
(T/'format_check.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(checks,ensure_ascii=False,indent=2))
