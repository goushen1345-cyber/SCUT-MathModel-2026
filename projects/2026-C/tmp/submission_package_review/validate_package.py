from pathlib import Path
import json,hashlib,ast,shutil,zipfile,io
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.packaging.core import DocumentProperties
from datetime import datetime
P=Path(__file__).resolve().parents[2];D=P/'支撑材料';T=Path(__file__).resolve().parent
report={}
for label,file in [('Q2','result2'),('Q3','result3'),('Q4-2','result4-2'),('Q4-3','result4-3')]:
 z=dict(np.load(D/'results'/f'{label}.npz'));w=openpyxl.load_workbook(D/f'{file}.xlsx',data_only=True);ex=z['execution']
 for sheet,order in [('计划购电量',z.get('q0',z.get('order'))),*([('调整购电量',z['versions'][:,-1])] if 'versions' in z else [])]:
  s=w[sheet];arr=np.array([[s.cell(i+2,145 if k==0 else k+1).value for k in range(144)] for i in range(334)],float)
  np.testing.assert_allclose(arr,order,atol=1e-9,rtol=0)
  for i in range(334):
   assert s.cell(i+2,1).value.strftime('%Y-%m-%d')==str((pd.Timestamp('2025-02-01')+pd.Timedelta(days=i)).date())
   assert abs(s.cell(i+2,146).value-order[i].sum())<1e-8
 s=w['充放电量'];assert s.max_row==2005
 for i in range(334):
  for g in range(6):
   r=2+i*6+g
   assert abs(s.cell(r,3).value-ex[i,g*24:(g+1)*24,5:7].sum())<1e-8
   assert abs(s.cell(r,4).value-ex[i,g*24:(g+1)*24,7].sum())<1e-8
  assert abs(s.cell(2+i*6,6).value-ex[i,0,0])<1e-8
  assert abs(s.cell(3+i*6,6).value-ex[i,-1,1])<1e-8
 actualsum=sum(float(row[2]) for row in w['紧急购电量'].iter_rows(min_row=2,values_only=True) if row[2] is not None)
 assert abs(actualsum-ex[:,:,8].sum())<1e-7
 report[file]=dict(orders_exact=True,dates_complete=True,charge_groups=2004,endpoints=668,emergency_total_exact=True)
for f in (D/'templates').glob('*.xlsx'):
 w=openpyxl.load_workbook(f);w.properties=DocumentProperties(creator='',lastModifiedBy='',created=datetime(2026,9,13),modified=datetime(2026,9,13));w.save(f)
for p in (D/'code').glob('*.py'):ast.parse(p.read_text('utf-8'),filename=p.name)
(D/'evidence/workbook_readback.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
# Move only our generated runtime artifacts out of the submission directory.
for name in ['reproduced','code/__pycache__','tmp']:
 src=(D/name).resolve();assert src.is_relative_to(D.resolve())
 if src.exists():
  dst=T/('checked_'+name.replace('/','_'));assert not dst.exists(),dst
  shutil.move(str(src),str(dst))
shutil.copy2(P/'paper/提交配套/AI工具使用详情.pdf',D/'AI工具使用详情.pdf')
bad=[]
for f in D.rglob('*'):
 if f.is_file() and f.suffix in ['.py','.md','.txt','.json','.csv','.tex']:
  s=f.read_text('utf-8-sig')
  for term in ['SCUT','Users/Rog','Users\\Rog','华南理工','wxid_']:
   if term in s:bad.append((str(f.relative_to(D)),term))
assert not bad,bad
buf=io.BytesIO()
with zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
 for p in D.rglob('*'):
  if p.is_file():z.write(p,str(Path('支撑材料')/p.relative_to(D)))
print(json.dumps(dict(readback=report,anonymous_scan='no listed identity markers',zip_estimate_bytes=len(buf.getbuffer()),file_count=sum(p.is_file() for p in D.rglob('*'))),ensure_ascii=False))
assert len(buf.getbuffer())<20_000_000
(T/'size_check.json').write_text(json.dumps(dict(estimated_zip_bytes=len(buf.getbuffer()),note='In-memory ZIP measurement only; no claim that WinRAR was used.'),indent=2),encoding='utf-8')
