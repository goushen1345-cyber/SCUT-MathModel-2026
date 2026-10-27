from pathlib import Path
import hashlib
import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt


HERE = Path(__file__).resolve().parent
PROJECT = HERE.parents[1]
PREP = PROJECT / "data/prepared-v1"
RAW = PROJECT / "data/raw/C题/附件/附件4.xlsx"
OUT = HERE / "outputs"
FIG = HERE / "figures"


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def savefig(fig, name):
    outputs = []
    for ext in ("pdf", "png"):
        path = FIG / f"{name}.{ext}"
        kwargs = {"metadata": {"CreationDate": None, "ModDate": None}} if ext == "pdf" else {}
        fig.savefig(path, dpi=300, bbox_inches="tight", **kwargs)
        outputs.append(path)
    plt.close(fig)
    return outputs


def longest_run(mask):
    best = current = 0
    for value in mask:
        current = current + 1 if value else 0
        best = max(best, current)
    return int(best)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    FIG.mkdir(parents=True, exist_ok=True)
    manifest = json.loads((PREP / "manifest.json").read_text(encoding="utf-8"))
    source_files = {
        "actual_price.csv": PREP / "actual_price.csv",
        "actual_power.csv": PREP / "actual_power.csv",
        "day_reference.csv": PREP / "day_reference.csv",
        "附件4.xlsx": RAW,
    }
    assert sha(source_files["actual_price.csv"]) == manifest["outputs"]["actual_price.csv"]["sha256"]

    price = pd.read_csv(source_files["actual_price.csv"], float_precision="round_trip")
    assert len(price) == 365 * 144
    price["date"] = pd.to_datetime(price["source_date"])
    price["slot_minute"] = price["source_minute"] % 1440
    price["datetime"] = price["date"] + pd.to_timedelta(price["slot_minute"], unit="m")
    price = price.sort_values(["date", "slot_minute"]).reset_index(drop=True)
    assert not price.duplicated(["date", "slot_minute"]).any()
    assert price.price_yuan_per_kwh.notna().all()
    assert (price.price_yuan_per_kwh > 0).all()
    assert price.groupby("date").size().eq(144).all()

    # Verify every prepared value against the raw workbook in source cell order.
    wb = load_workbook(RAW, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    assert ws.max_row == 366 and ws.max_column == 145
    raw_values = np.array(
        [list(row) for row in ws.iter_rows(
            min_row=2, max_row=366, min_col=2, max_col=145, values_only=True
        )],
        dtype=float,
    )
    wb.close()
    source_order = pd.read_csv(source_files["actual_price.csv"], float_precision="round_trip")
    prepared_values = source_order.price_yuan_per_kwh.to_numpy().reshape(365, 144)
    raw_max_diff = float(np.max(np.abs(raw_values - prepared_values)))
    assert raw_max_diff < 1e-12

    p = price.price_yuan_per_kwh
    price["month"] = price.date.dt.month
    price["weekday"] = price.date.dt.dayofweek
    price["hour"] = price.slot_minute / 60
    price["day_index"] = (price.date - pd.Timestamp("2025-01-01")).dt.days
    weekday_names = {0: "周一", 1: "周二", 2: "周三", 3: "周四", 4: "周五", 5: "周六", 6: "周日"}

    profile = price.groupby("slot_minute").price_yuan_per_kwh.agg(
        mean="mean", median="median", std="std",
        p10=lambda x: x.quantile(.10), p90=lambda x: x.quantile(.90),
        p95=lambda x: x.quantile(.95), minimum="min", maximum="max"
    ).reset_index()
    profile["hour"] = profile.slot_minute / 60

    monthly = price.groupby("month").price_yuan_per_kwh.agg(
        mean="mean", median="median", std="std",
        p90=lambda x: x.quantile(.90), p95=lambda x: x.quantile(.95),
        minimum="min", maximum="max"
    ).reset_index()
    weekday = price.groupby("weekday").price_yuan_per_kwh.agg(
        mean="mean", median="median", std="std",
        p90=lambda x: x.quantile(.90)
    ).reset_index()
    weekday["label"] = weekday.weekday.map(weekday_names)

    daily = price.groupby("date").price_yuan_per_kwh.agg(
        mean="mean", median="median", std="std", minimum="min", maximum="max"
    ).reset_index()
    daily["range"] = daily.maximum - daily.minimum
    daily["month"] = daily.date.dt.month
    daily["weekday"] = daily.date.dt.dayofweek

    bands = [
        (0, 6, "00--06时"), (6, 10, "06--10时"), (10, 16, "10--16时"),
        (16, 19, "16--19时"), (19, 22, "19--22时"), (22, 24, "22--24时"),
    ]
    band_rows = []
    for start, end, label in bands:
        z = price[(price.hour >= start) & (price.hour < end)].price_yuan_per_kwh
        band_rows.append({
            "band": label, "mean": z.mean(), "median": z.median(),
            "p90": z.quantile(.90), "std": z.std(),
        })
    band = pd.DataFrame(band_rows)

    ordered = price.price_yuan_per_kwh.reset_index(drop=True)
    lags = [1, 6, 36, 144, 288, 1008]
    lag_labels = ["10分钟", "1小时", "6小时", "前1天同刻", "前2天同刻", "前7天同刻"]
    autocorr = pd.DataFrame({
        "lag": lags,
        "label": lag_labels,
        "correlation": [ordered.corr(ordered.shift(lag)) for lag in lags],
        "mae": [float((ordered - ordered.shift(lag)).abs().dropna().mean()) for lag in lags],
    })
    change = ordered.diff().dropna()
    abs_change = change.abs()
    change_stats = {
        "median_abs_10min": float(abs_change.median()),
        "p90_abs_10min": float(abs_change.quantile(.90)),
        "p95_abs_10min": float(abs_change.quantile(.95)),
        "p99_abs_10min": float(abs_change.quantile(.99)),
        "max_abs_10min": float(abs_change.max()),
    }

    pivot = price.pivot(index="date", columns="slot_minute", values="price_yuan_per_kwh")
    mean_curve = pivot.mean(axis=0)
    curve_corr = pivot.apply(lambda row: row.corr(mean_curve), axis=1)
    previous_day_mae = (pivot.iloc[1:].to_numpy() - pivot.iloc[:-1].to_numpy())
    previous_day_mae = np.mean(np.abs(previous_day_mae), axis=1)

    q95, q99 = float(p.quantile(.95)), float(p.quantile(.99))
    price["spike95"] = price.price_yuan_per_kwh >= q95
    price["spike99"] = price.price_yuan_per_kwh >= q99
    spike_hour = price.groupby("slot_minute").agg(
        spike95_rate=("spike95", "mean"), spike99_rate=("spike99", "mean")
    ).reset_index()
    spike_hour["hour"] = spike_hour.slot_minute / 60
    spike_month = price.groupby("month").agg(
        spike95_count=("spike95", "sum"), spike99_count=("spike99", "sum")
    ).reset_index()

    power = pd.read_csv(source_files["actual_power.csv"], float_precision="round_trip")
    power["date"] = pd.to_datetime(power.source_date)
    power["slot_minute"] = power.source_minute % 1440
    merged = price.merge(
        power[["date", "slot_minute", "load_kw", "pv_actual_kw"]],
        on=["date", "slot_minute"], how="inner", validate="one_to_one"
    )
    assert len(merged) == len(price)
    merged["net_load_kw"] = merged.load_kw - merged.pv_actual_kw
    overall_corr = {
        "load": float(merged.price_yuan_per_kwh.corr(merged.load_kw)),
        "pv": float(merged.price_yuan_per_kwh.corr(merged.pv_actual_kw)),
        "net_load": float(merged.price_yuan_per_kwh.corr(merged.net_load_kw)),
    }
    for col in ["price_yuan_per_kwh", "load_kw", "pv_actual_kw", "net_load_kw"]:
        merged[col + "_resid"] = merged[col] - merged.groupby("slot_minute")[col].transform("mean")
    residual_corr = {
        "load": float(merged.price_yuan_per_kwh_resid.corr(merged.load_kw_resid)),
        "pv": float(merged.price_yuan_per_kwh_resid.corr(merged.pv_actual_kw_resid)),
        "net_load": float(merged.price_yuan_per_kwh_resid.corr(merged.net_load_kw_resid)),
    }
    merged["net_resid_decile"] = pd.qcut(merged.net_load_kw_resid, 10, labels=False, duplicates="drop") + 1
    decile = merged.groupby("net_resid_decile").agg(
        net_load_residual_kw=("net_load_kw_resid", "mean"),
        price_residual=("price_yuan_per_kwh_resid", "mean"),
        count=("price_yuan_per_kwh", "size")
    ).reset_index()

    ref = pd.read_csv(source_files["day_reference.csv"], float_precision="round_trip")
    ref["slot_minute"] = ref.source_minute % 1440
    ref = ref.sort_values("slot_minute")
    fixed_price = ref.price_yuan_per_kwh.to_numpy()
    assert len(fixed_price) == 144
    fixed_corr = float(np.corrcoef(profile["mean"], fixed_price)[0, 1])
    fixed_max_abs_diff = float(np.max(np.abs(profile["mean"].to_numpy() - fixed_price)))

    peak_slot = int(profile.loc[profile["mean"].idxmax(), "slot_minute"])
    valley_slot = int(profile.loc[profile["mean"].idxmin(), "slot_minute"])
    daily_peak_slot = pivot.idxmax(axis=1)
    daily_valley_slot = pivot.idxmin(axis=1)
    peak_window_share = float(((daily_peak_slot >= 19 * 60) & (daily_peak_slot < 22 * 60)).mean())
    valley_window_share = float(((daily_valley_slot >= 10 * 60) & (daily_valley_slot < 16 * 60)).mean())
    valley_night_share = float(((daily_valley_slot < 6 * 60) | (daily_valley_slot >= 22 * 60)).mean())

    overall = {
        "rows": int(len(price)),
        "days": int(price.date.nunique()),
        "intervals_per_day": 144,
        "minimum": float(p.min()),
        "maximum": float(p.max()),
        "mean": float(p.mean()),
        "median": float(p.median()),
        "std": float(p.std()),
        "cv": float(p.std() / p.mean()),
        "p05": float(p.quantile(.05)),
        "p25": float(p.quantile(.25)),
        "p75": float(p.quantile(.75)),
        "p95": q95,
        "p99": q99,
        "mean_profile_peak_slot_minute": peak_slot,
        "mean_profile_peak_price": float(profile.loc[profile["mean"].idxmax(), "mean"]),
        "mean_profile_valley_slot_minute": valley_slot,
        "mean_profile_valley_price": float(profile.loc[profile["mean"].idxmin(), "mean"]),
        "daily_peak_in_19_22_share": peak_window_share,
        "daily_valley_in_10_16_share": valley_window_share,
        "daily_valley_in_22_06_share": valley_night_share,
        "daily_curve_correlation_mean": float(curve_corr.mean()),
        "daily_curve_correlation_median": float(curve_corr.median()),
        "daily_curve_correlation_p10": float(curve_corr.quantile(.10)),
        "previous_day_same_slot_mae_mean": float(np.mean(previous_day_mae)),
        "previous_day_same_slot_mae_median": float(np.median(previous_day_mae)),
        "fixed_vs_actual_mean_profile_correlation": fixed_corr,
        "fixed_vs_actual_mean_profile_max_abs_diff": fixed_max_abs_diff,
        "raw_prepared_max_abs_diff": raw_max_diff,
        "longest_consecutive_top5pct_intervals": longest_run(price.spike95.to_numpy()),
    }
    summary = {
        "scope": "descriptive exploration only; no forecast fitting or dispatch optimization",
        "time_mapping": "source_minute modulo 1440; source 0:00+1 is mapped to the same row date 00:00--00:10 under the confirmed project convention",
        "source_sha256": {name: sha(path) for name, path in source_files.items()},
        "overall": overall,
        "change": change_stats,
        "autocorrelation": dict(zip(lag_labels, autocorr.correlation.astype(float))),
        "same_series_lag_mae": dict(zip(lag_labels, autocorr.mae.astype(float))),
        "correlation_with_actual_supply": {
            "overall": overall_corr,
            "after_removing_time_of_day_mean": residual_corr,
            "boundary": "descriptive association only; future actual load/PV cannot be used as predictors",
        },
        "monthly": monthly.to_dict("records"),
        "weekday": weekday.to_dict("records"),
        "bands": band.to_dict("records"),
    }

    profile.to_csv(OUT / "intraday_profile.csv", index=False)
    monthly.to_csv(OUT / "monthly.csv", index=False)
    weekday.to_csv(OUT / "weekday.csv", index=False)
    daily.to_csv(OUT / "daily.csv", index=False)
    band.to_csv(OUT / "time_bands.csv", index=False)
    autocorr.to_csv(OUT / "autocorrelation.csv", index=False)
    spike_hour.to_csv(OUT / "spike_by_time.csv", index=False)
    spike_month.to_csv(OUT / "spike_by_month.csv", index=False)
    decile.to_csv(OUT / "net_load_residual_deciles.csv", index=False)
    (OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

    plt.rcParams.update({
        "font.family": "SimSun", "font.size": 10, "axes.unicode_minus": False,
        "text.color": "black", "axes.labelcolor": "black", "xtick.color": "black",
        "ytick.color": "black", "pdf.fonttype": 42, "axes.spines.top": False,
        "axes.spines.right": False, "axes.axisbelow": True, "savefig.facecolor": "white"
    })
    blue, green, orange, red, grey = "#3f7198", "#648b6b", "#c58d3f", "#b54f4b", "#777777"
    generated = []

    fig, axes = plt.subplots(2, 2, figsize=(8.2, 5.7))
    x = profile.hour.to_numpy()
    axes[0, 0].fill_between(x, profile.p10.to_numpy(), profile.p90.to_numpy(), color=blue, alpha=.16, label="10%—90%区间")
    axes[0, 0].plot(x, profile["mean"], color=blue, linewidth=1.5, label="全年同刻均值")
    axes[0, 0].plot(x, fixed_price, color=grey, linestyle="--", linewidth=1.0, label="附件1固定电价")
    axes[0, 0].set(xlabel="时刻 / h", ylabel="电价 / 元·kWh$^{-1}$", xticks=np.arange(0, 25, 4), xlim=(0, 24))
    axes[0, 0].legend(frameon=False, fontsize=8)
    axes[0, 0].set_title("(a) 日内形状稳定，但每天高低不同")

    hourly = price.assign(hour_int=(price.slot_minute // 60).astype(int)).groupby(["day_index", "hour_int"]).price_yuan_per_kwh.mean().unstack()
    im = axes[0, 1].imshow(hourly.to_numpy(), aspect="auto", origin="lower", cmap="YlOrRd",
                           extent=[0, 24, 1, 365], vmin=float(p.quantile(.05)), vmax=float(p.quantile(.95)))
    axes[0, 1].set(xlabel="时刻 / h", ylabel="日期序号", xticks=np.arange(0, 25, 4))
    axes[0, 1].set_title("(b) 全年逐日逐时电价")
    cb = fig.colorbar(im, ax=axes[0, 1], fraction=.046, pad=.04)
    cb.set_label("元/kWh")

    axes[1, 0].plot(monthly.month, monthly["mean"], color=blue, marker="o", label="均值")
    axes[1, 0].plot(monthly.month, monthly.p90, color=orange, marker="s", label="90%分位")
    axes[1, 0].set(xlabel="月份", ylabel="电价 / 元·kWh$^{-1}$", xticks=monthly.month)
    axes[1, 0].set_title("(c) 月份主要改变价格水平与高价尾部")
    axes[1, 0].legend(frameon=False)

    axes[1, 1].bar(weekday.label, weekday["mean"], color=blue, width=.62, label="均值")
    axes[1, 1].plot(weekday.label, weekday.p90, color=orange, marker="o", label="90%分位")
    axes[1, 1].set(xlabel="星期", ylabel="电价 / 元·kWh$^{-1}$")
    axes[1, 1].set_title("(d) 星期差异")
    axes[1, 1].legend(frameon=False)
    for ax in axes.ravel():
        ax.grid(axis="y", alpha=.17)
    fig.suptitle("附件4波动电价的日内、月份与星期结构", fontsize=13)
    fig.tight_layout()
    generated += savefig(fig, "01_price_structure")

    fig, axes = plt.subplots(2, 2, figsize=(8.2, 5.7))
    axes[0, 0].bar(autocorr.label, autocorr.correlation, color=[blue, blue, blue, green, green, green])
    axes[0, 0].set(ylim=(0, 1), ylabel="相关系数")
    axes[0, 0].tick_params(axis="x", rotation=25)
    axes[0, 0].set_title("(a) 相邻价格和同刻历史都有参考价值")

    quantile_levels = [.50, .90, .95, .99]
    qvals = [abs_change.quantile(q) for q in quantile_levels]
    axes[0, 1].bar(["中位数", "90%分位", "95%分位", "99%分位"], qvals, color=[blue, blue, orange, red])
    axes[0, 1].set(ylabel="相邻10分钟绝对变化 / 元·kWh$^{-1}$")
    axes[0, 1].set_title("(b) 多数变化较小，尾部跳动明显")

    axes[1, 0].plot(spike_hour.hour, 100 * spike_hour.spike95_rate, color=red, linewidth=1.4)
    axes[1, 0].fill_between(spike_hour.hour, 0, 100 * spike_hour.spike95_rate, color=red, alpha=.15)
    axes[1, 0].set(xlabel="时刻 / h", ylabel="进入全年最高5%的比例 / %", xticks=np.arange(0, 25, 4), xlim=(0, 24))
    axes[1, 0].set_title("(c) 高价风险集中在特定时段")

    axes[1, 1].plot(decile.net_resid_decile, decile.price_residual, color=green, marker="o")
    axes[1, 1].axhline(0, color=grey, linewidth=.8)
    axes[1, 1].set(xlabel="去除日内均值后的净负载分组", ylabel="同刻电价偏差均值 / 元·kWh$^{-1}$", xticks=decile.net_resid_decile)
    axes[1, 1].set_title("(d) 同一时刻净负载偏高时，价格往往也偏高")
    for ax in axes.ravel():
        ax.grid(axis="y", alpha=.17)
    fig.suptitle("波动电价的可预测线索与高价风险", fontsize=13)
    fig.tight_layout()
    generated += savefig(fig, "02_price_predictability")

    figure_manifest = {
        "scope": "descriptive plots from verified attachment 4 and prepared-v1; no forecasting or optimization",
        "script_sha256": sha(__file__),
        "source_sha256": {name: sha(path) for name, path in source_files.items()},
        "matplotlib": matplotlib.__version__,
        "font": "SimSun",
        "dpi": 300,
        "outputs": {str(path.relative_to(HERE)): sha(path) for path in generated},
    }
    (FIG / "manifest.json").write_text(json.dumps(figure_manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "summary": str(OUT / "summary.json"),
        "figures": [str(path) for path in generated],
        "peak_slot": peak_slot,
        "valley_slot": valley_slot,
        "raw_max_diff": raw_max_diff,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
