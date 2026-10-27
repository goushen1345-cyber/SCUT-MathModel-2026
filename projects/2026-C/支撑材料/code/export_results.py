"""Export saved final results into the five official workbooks and check balances."""
from pathlib import Path
from datetime import datetime,timedelta
from copy import copy,deepcopy
import csv,gzip,json
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.packaging.core import DocumentProperties
ROOT=Path(__file__).resolve().parents[1]
EXPECTED={'Q1':35126.85,'Q2':13583810.84,'Q3':13014172.33,'Q4-2':14248570.87,'Q4-3':13695929.70}
EXEC_NAMES=['storage_start_kwh','storage_end_kwh','received_kwh','grid_load_kwh','pv_load_kwh','grid_charge_kwh','pv_charge_kwh','discharge_kwh','emergency_kwh','unused_order_kwh','unused_pv_kwh','load_kwh']
def time_label(k):return f'{k//6:02d}:{k%6*10:02d}'
def metadata(w,template_name):
 w.properties=DocumentProperties(creator='',lastModifiedBy='',title='',description='',created=datetime(2026,9,13),modified=datetime(2026,9,13))
 for s in w:
  for row in s:
   for c in row:
    c.comment=None
    if isinstance(c.value,(int,float)):c.number_format='0.00'
 template=openpyxl.load_workbook(ROOT/'templates'/template_name)
 format_workbook(w,template)
 template.close()

def format_workbook(w,template):
 """Restore template styles, extending daily blocks without changing any values."""
 cache={}
 for s in w:
  t=template[s.title]
  for attr in ['sheet_format','sheet_properties','views','auto_filter',
               'page_setup','page_margins','print_options']:
   setattr(s,attr,deepcopy(getattr(t,attr)))
  s.print_title_rows=t.print_title_rows
  s.print_title_cols=t.print_title_cols
  s.print_area=str(t.print_area) if t.print_area else ''
  s.row_dimensions.clear();s.column_dimensions.clear()
  for key,dim in t.column_dimensions.items():
   s.column_dimensions[key]=copy(dim);s.column_dimensions[key].parent=s
  for key,dim in t.row_dimensions.items():
   s.row_dimensions[key]=copy(dim);s.row_dimensions[key].parent=s
  annual_storage=s.title=='充放电量' and s.max_column==6
  for row in s:
   r=row[0].row;single=False
   if r==1:tr=1
   elif annual_storage:tr=2+(r-2)%6
   elif s.title=='紧急购电量':
    date=s.cell(r,1).value
    first=r==2 or s.cell(r-1,1).value!=date
    last=r==s.max_row or s.cell(r+1,1).value!=date
    tr=2 if first else (4 if last else 3)
    single=first and last
   else:tr=r if r<=t.max_row else 2
   if tr in t.row_dimensions:
    s.row_dimensions[r]=copy(t.row_dimensions[tr])
    s.row_dimensions[r].index=r;s.row_dimensions[r].parent=s
   for c in row:
    ref=t.cell(tr,c.column)
    # All retained dates use the template's date format, including continuation rows.
    number_format=t.cell(2,1).number_format if isinstance(c.value,datetime) else ref.number_format
    key=(s.title,ref.style_id,number_format,single,c.column if single else 0)
    if key not in cache:
     for attr in ['font','fill','border','alignment','protection']:
      setattr(c,attr,copy(getattr(ref,attr)))
     c.number_format=number_format
     if single:
      border=copy(c.border);border.bottom=copy(t.cell(4,c.column).border.bottom)
      c.border=border
     cache[key]=copy(c._style)
    c._style=copy(cache[key])
def numeric(s,r,c,v):
 if abs(v)<1e-10:v=0.
 s.cell(r,c,float(v)).number_format='0.00'
def clear_data(s):
 for row in s.iter_rows(min_row=2):
  for c in row:c.value=None
def fill_orders(s,orders,fee):
 for i,order in enumerate(orders):
  r=i+2
  for k,v in enumerate(order):numeric(s,r,145 if k==0 else k+1,v)
  numeric(s,r,146,order.sum());numeric(s,r,147,fee[i])
 s.freeze_panes='B2'
 # The official template's last header is retained; its same-date slot-0 meaning is documented.
def fill_storage(s,ex):
 styles=[copy(c._style) for c in s[2]];clear_data(s)
 for i,day in enumerate(ex):
  date=datetime(2025,2,1)+timedelta(days=i)
  for g in range(6):
   r=2+i*6+g
   for j,style in enumerate(styles,1):s.cell(r,j)._style=copy(style)
   s.cell(r,1,date).number_format='yyyy-mm-dd'
   s.cell(r,2,f'{4*g}:00-{4*(g+1)}:00')
   numeric(s,r,3,day[g*24:(g+1)*24,5:7].sum());numeric(s,r,4,day[g*24:(g+1)*24,7].sum())
   if g<2:s.cell(r,5,'0:00' if g==0 else '24:00');numeric(s,r,6,day[0,0] if g==0 else day[-1,1])
 s.freeze_panes='C2'
def fill_emergency(s,ex):
 clear_data(s);r=2;events=0
 for i,day in enumerate(ex):
  k=0;found=False;date=datetime(2025,2,1)+timedelta(days=i)
  while k<144:
   if day[k,8]<=1e-8:k+=1;continue
   end=k+1
   while end<144 and day[end,8]>1e-8:end+=1
   s.cell(r,1,date).number_format='yyyy-mm-dd';s.cell(r,2,f'{time_label(k)}-{time_label(end)}');numeric(s,r,3,day[k:end,8].sum())
   k=end;r+=1;events+=1;found=True
  if not found:s.cell(r,1,date).number_format='yyyy-mm-dd';s.cell(r,2,'无');numeric(s,r,3,0);r+=1
 s.column_dimensions['A'].width=15;s.column_dimensions['B'].width=22;s.column_dimensions['C'].width=18
 return events
def main():
 src=np.load(ROOT/'data/inputs.npz');report={};dailyrows=[]
 q1=pd.read_csv(ROOT/'results/Q1.csv',float_precision='round_trip')
 w=openpyxl.load_workbook(ROOT/'templates/result1.xlsx')
 for k,v in enumerate(q1.purchase_kwh):numeric(w['计划购电量'],145 if k==0 else k+1,2,v)
 for g in range(6):
  part=q1.iloc[g*24:(g+1)*24];numeric(w['充放电量'],g+2,2,part.charge_kwh.sum());numeric(w['充放电量'],g+2,3,part.discharge.sum())
 numeric(w['充放电量'],2,5,q1.state_start.iloc[0]);numeric(w['充放电量'],3,5,q1.state_end.iloc[-1]);metadata(w,'result1.xlsx');w.save(ROOT/'result1.xlsx')
 report['Q1']={'total_cost':float(q1.cost_yuan.sum()),'order_kwh':float(q1.purchase_kwh.sum())}
 for label in ['Q2','Q3','Q4-2','Q4-3']:
  z=dict(np.load(ROOT/'results'/f'{label}.npz'));ex=z['execution'];rolling='versions' in z
  order=z['versions'][:,-1,:] if rolling else z['order'];zero=z['q0'] if rolling else order
  price=src['actual_prices'][31:365] if label.startswith('Q4') else np.tile(src['price'],(334,1))
  normal=(price*order).sum(1);adjust=(.5*price*abs(order-zero)).sum(1);emergency=(5*price*ex[:,:,8]).sum(1)
  cost=normal+adjust+emergency;assert abs(cost.sum()-EXPECTED[label])<.005,(label,cost.sum())
  balance=ex[:,:,3]+ex[:,:,4]+ex[:,:,7]+ex[:,:,8]-ex[:,:,11]
  storage=ex[:,:,1]-ex[:,:,0]-.9*(ex[:,:,5]+ex[:,:,6])+ex[:,:,7]/.9
  assert max(abs(balance).max(),abs(storage).max())<1e-6
  assert ex[:,:,:2].min()>=1200-1e-6 and ex[:,:,:2].max()<=10800+1e-6
  assert max((ex[:,:,5]+ex[:,:,6]).max(),ex[:,:,7].max())<=5000/6+1e-6
  assert np.max(abs(ex[1:,0,0]-ex[:-1,-1,1]))<1e-6
  if rolling:
   for a,r in enumerate([36,72,108],1):assert np.max(abs(z['versions'][:,a,:r]-z['versions'][:,a-1,:r]))<1e-7
  file={'Q2':'result2','Q3':'result3','Q4-2':'result4-2','Q4-3':'result4-3'}[label]
  w=openpyxl.load_workbook(ROOT/'templates'/f'{file}.xlsx');fill_orders(w['计划购电量'],zero,(price*zero).sum(1))
  if rolling:fill_orders(w['调整购电量'],order,normal+adjust)
  fill_storage(w['充放电量'],ex);events=fill_emergency(w['紧急购电量'],ex);metadata(w,f'{file}.xlsx');w.save(ROOT/f'{file}.xlsx')
  if rolling:
   with gzip.open(ROOT/'results'/f'{label}_adjustments.csv.gz','wt',encoding='utf-8-sig',newline='') as f:
    writer=csv.writer(f);writer.writerow(['date','release_time','interval_start','interval_end','midnight_kwh','revised_total_kwh','net_change_vs_midnight_kwh'])
    for i in range(334):
     date=str((datetime(2025,2,1)+timedelta(days=i)).date())
     for a,r in enumerate([0,36,72,108]):
      for k in range(r,144):writer.writerow([date,time_label(r),time_label(k),time_label(k+1),zero[i,k],z['versions'][i,a,k],z['versions'][i,a,k]-zero[i,k]])
  for i in range(334):dailyrows.append(dict(question=label,date=str((datetime(2025,2,1)+timedelta(days=i)).date()),midnight_order_kwh=float(zero[i].sum()),final_order_kwh=float(order[i].sum()),normal_cost=normal[i],adjustment_cost=adjust[i],emergency_cost=emergency[i],total_cost=cost[i],emergency_kwh=float(ex[i,:,8].sum()),initial_kwh=ex[i,0,0],final_kwh=ex[i,-1,1]))
  report[label]=dict(days=334,slots=48096,storage_groups=2004,emergency_events=events,total_cost=float(cost.sum()),normal_cost=float(normal.sum()),adjustment_cost=float(adjust.sum()),emergency_cost=float(emergency.sum()),energy_balance_max_error=float(abs(balance).max()),storage_max_error=float(abs(storage).max()))
 pd.DataFrame(dailyrows).to_csv(ROOT/'results/daily_costs.csv',index=False,encoding='utf-8-sig',float_format='%.17g')
 (ROOT/'evidence/export_check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
 print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
