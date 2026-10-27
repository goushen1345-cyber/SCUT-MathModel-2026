"""Read original G0 sources; preserve hashes, PDF text and structural evidence only."""
from datetime import date, datetime, time
import hashlib
import json
from pathlib import Path
import sys

import openpyxl
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw/C题"
OUT = ROOT / "tmp/problem-analysis"


def plain(value):
    if isinstance(value, (date, datetime, time)):
        return {"type": type(value).__name__, "value": value.isoformat()}
    return value


def run():
    OUT.mkdir(parents=True, exist_ok=True)
    all_files = sorted(p for p in RAW.rglob("*") if p.is_file())
    files = [p for p in all_files if not p.name.startswith("~$")]
    hashes = {p.relative_to(ROOT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
              for p in files}
    snapshot = {"scope": "G0 source text, labels and boundary cells; no data-quality audit or processing",
                "python": sys.version, "openpyxl": openpyxl.__version__,
                "sha256": hashes, "excluded_lock_files": [p.relative_to(ROOT).as_posix()
                    for p in all_files if p.name.startswith("~$")], "workbooks": {}}
    for p in files:
        if p.suffix.lower() != ".xlsx":
            continue
        wb = openpyxl.load_workbook(p, read_only=True, data_only=False)
        sheets = {}
        for ws in wb:
            cells = {}
            for row in ws.iter_rows(min_row=1, max_row=min(6, ws.max_row)):
                for cell in row:
                    if cell.value is not None:
                        cells[cell.coordinate] = plain(cell.value)
            for row in ws.iter_rows(min_row=max(1, ws.max_row-1), max_row=ws.max_row, max_col=3):
                for cell in row:
                    if cell.value is not None:
                        cells[cell.coordinate] = plain(cell.value)
            sheets[ws.title] = {"dimension": ws.calculate_dimension(),
                                "cells": cells}
        snapshot["workbooks"][p.relative_to(ROOT).as_posix()] = sheets
        wb.close()
    pdf = PdfReader(RAW / "C题.pdf")
    snapshot["pdf_pages"] = len(pdf.pages)
    (OUT / "source_text.txt").write_text("\n\n".join(
        f"PDF PAGE {i + 1}\n{page.extract_text()}" for i, page in enumerate(pdf.pages)), encoding="utf-8")
    (OUT / "source_snapshot.json").write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8")
    after = {p.relative_to(ROOT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
    if hashes != after:
        raise RuntimeError("Original sources changed during read-only inspection")
    print(json.dumps({"pdf_pages": len(pdf.pages), "workbooks": len(snapshot["workbooks"]),
                      "lock_files_excluded": len(snapshot["excluded_lock_files"]),
                      "source_hashes_unchanged": True, "pdf_sha256": hashes["data/raw/C题/C题.pdf"]}, ensure_ascii=False))


if __name__ == "__main__":
    run()
