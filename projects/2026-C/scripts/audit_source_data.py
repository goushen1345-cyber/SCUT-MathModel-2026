"""Read-only quality audit. No cleaning, fitting, forecasting or workbook export.

Run from repository root with bundled Python:
  python -X utf8 projects/2026-C/scripts/audit_source_data.py
Outputs metadata/counts/source-cell examples to audit/data_audit_v1.json.
Full-year ranges are quality checks only, not predictive-model calibration.
"""
from collections import Counter
from datetime import date, datetime, time, timedelta
import hashlib
import json
import math
from pathlib import Path
import re
import sys

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / 'data/raw/C题'
OUT = ROOT / 'audit/data_audit_v1.json'
EXAMPLE_LIMIT = 12


def plain(v):
    if isinstance(v, (date, datetime, time)):
        return {'type': type(v).__name__, 'value': v.isoformat()}
    return v


def cell_address(r, c):
    return f'{get_column_letter(c)}{r}'


def day(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        return date.fromisoformat('-'.join(f'{int(x):02d}' if i else x
            for i, x in enumerate(re.split(r'[-/.]', v.strip()))))
    raise ValueError(f'Unknown date type: {type(v).__name__}')


def minute(v):
    if isinstance(v, (time, datetime)):
        if v.second or v.microsecond:
            raise ValueError('Non-minute time')
        return v.hour * 60 + v.minute
    if isinstance(v, str):
        match = re.fullmatch(r'(\d{1,2}):(\d{1,2})(\+1)?', v.strip())
        if match:
            h, m = int(match[1]), int(match[2])
            if 0 <= h <= 24 and 0 <= m < 60 and (h < 24 or m == 0):
                return h * 60 + m + (1440 if match[3] else 0)
    raise ValueError(f'Unknown time: {v!r}')


def axis_dates(values, start, end):
    parsed, failures = [], []
    for r, value in values:
        try:
            parsed.append(day(value))
        except (ValueError, TypeError):
            failures.append({'cell': f'A{r}', 'value': plain(value)})
    expected = [start + timedelta(days=k) for k in range((end-start).days+1)]
    counts = Counter(parsed)
    return {'count': len(values), 'type_counts': dict(Counter(type(v).__name__ for _, v in values)),
            'parse_failures': failures, 'matches_expected_order': parsed == expected,
            'missing_dates': [v.isoformat() for v in expected if v not in counts],
            'extra_dates': [v.isoformat() for v in counts if v not in set(expected)],
            'duplicate_dates': {v.isoformat(): n for v, n in counts.items() if n > 1}}


def axis_times(values):
    parsed, failures = [], []
    for address, value in values:
        try:
            parsed.append(minute(value))
        except ValueError:
            failures.append({'cell': address, 'value': plain(value)})
    return {'count': len(values), 'type_counts': dict(Counter(type(v).__name__ for _, v in values)),
            'parse_failures': failures, 'matches_10_to_1440_every_10': parsed == list(range(10, 1441, 10)),
            'duplicate_minutes': {str(v): n for v, n in Counter(parsed).items() if n > 1},
            'first': plain(values[0][1]), 'last': plain(values[-1][1]), 'minutes': parsed}


def numeric_region(rows, r0, r1, c0, c1, unit):
    counts = Counter()
    examples = {}
    minimum = maximum = None
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            value = rows[r-1][c-1]
            address = cell_address(r, c)
            if value is None or (isinstance(value, str) and not value.strip()):
                kind = 'blank'
            elif isinstance(value, bool):
                kind = 'boolean'
            elif isinstance(value, (int, float)):
                if not math.isfinite(value):
                    kind = 'nonfinite'
                else:
                    kind = 'numeric'
                    counts['negative'] += value < 0
                    counts['zero'] += value == 0
                    counts['positive'] += value > 0
                    if value < 0 and len(examples.setdefault('negative', [])) < EXAMPLE_LIMIT:
                        examples['negative'].append({'cell': address, 'value': value})
                    if minimum is None or value < minimum['value']:
                        minimum = {'cell': address, 'value': value}
                    if maximum is None or value > maximum['value']:
                        maximum = {'cell': address, 'value': value}
            elif isinstance(value, str) and value.startswith('='):
                kind = 'formula'
            elif isinstance(value, str) and value.startswith('#'):
                kind = 'excel_error_or_marker'
            else:
                kind = 'nonnumeric'
            counts[kind] += 1
            if kind != 'numeric' and len(examples.setdefault(kind, [])) < EXAMPLE_LIMIT:
                examples[kind].append({'cell': address, 'value': plain(value)})
    return {'range': f'{cell_address(r0,c0)}:{cell_address(r1,c1)}', 'unit': unit,
            'expected_cells': (r1-r0+1)*(c1-c0+1), 'counts': dict(counts),
            'min': minimum, 'max': maximum, 'issue_examples': {k:v for k,v in examples.items() if v}}


def main():
    files = sorted(p for p in RAW.rglob('*') if p.is_file() and not p.name.startswith('~$'))
    before = {p.relative_to(ROOT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
    prior = json.loads((ROOT/'audit/g0_structure_v2.json').read_text(encoding='utf-8-sig'))
    changed = [k for k in set(before) | set(prior['sha256']) if before.get(k) != prior['sha256'].get(k)]
    if changed:
        raise RuntimeError(f'Source version changed: {changed}')
    result = {'recorded_at': datetime.now().astimezone().isoformat(), 'python': sys.version,
              'openpyxl': openpyxl.__version__, 'scope': 'read-only raw-data quality and layout audit; no transformation, EDA for model selection, fitting or workbook writing',
              'source_sha256': before, 'source_version_matches_prior': True,
              'script_sha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
              'excluded_locks': [p.relative_to(ROOT).as_posix() for p in RAW.rglob('~$*')],
              'data': {}, 'templates': {}}
    matrices = {}
    for path in (p for p in files if p.suffix.lower() == '.xlsx'):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        for ws in wb:
            rows = list(ws.iter_rows(values_only=True))
            key = f'{path.name}/{ws.title}'
            meta = {'source': path.relative_to(ROOT).as_posix(), 'sheet': ws.title,
                    'dimension': ws.calculate_dimension(), 'hidden': ws.sheet_state != 'visible'}
            if path.name == '附件1.xlsx':
                meta['time_axis'] = axis_times([(f'A{r}', rows[r-1][0]) for r in range(2,len(rows)+1)])
                meta['regions'] = {name: numeric_region(rows,2,len(rows),c,c,unit)
                    for name,c,unit in [('price',2,'元/kWh'),('load',3,'kW'),('pv_forecast',4,'kW')]}
            elif path.name in {'附件2.xlsx','附件4.xlsx'}:
                meta['date_axis'] = axis_dates([(r, rows[r-1][0]) for r in range(2,len(rows)+1)], date(2025,1,1),date(2025,12,31))
                meta['time_axis'] = axis_times([(cell_address(1,c),rows[0][c-1]) for c in range(2,len(rows[0])+1)])
                meta['regions'] = {'values': numeric_region(rows,2,len(rows),2,len(rows[0]),'元/kWh' if path.name == '附件4.xlsx' else 'kW')}
                signatures = Counter(tuple(row[1:]) for row in rows[1:])
                meta['identical_day_curve_groups'] = sum(n>1 for n in signatures.values())
                meta['identical_curve_extra_rows'] = sum(n-1 for n in signatures.values() if n>1)
                meta['all_zero_day_rows'] = [r for r in range(2,len(rows)+1) if all(v==0 for v in rows[r-1][1:])]
                matrices[key] = ([day(row[0]).isoformat() for row in rows[1:]],meta['time_axis']['minutes'])
            elif path.name == '附件3.xlsx':
                blocks = list(range(2,len(rows)+1,4))
                meta['date_axis'] = axis_dates([(r,rows[r-1][0]) for r in blocks],date(2025,1,1),date(2025,12,31))
                meta['group_date_blanks'] = sum(rows[r-1][0] in (None,'') for r in range(2,len(rows)+1) if r not in blocks)
                issues, beyond_year = [], 0
                forecast_keys = []
                for r in blocks:
                    d = day(rows[r-1][0])
                    for offset, release_hour in enumerate([0,6,12,18]):
                        rr = r+offset
                        if minute(rows[rr-1][1]) != release_hour*60:
                            issues.append({'cell': f'B{rr}', 'value': plain(rows[rr-1][1])})
                        release = datetime.combine(d,time(release_hour))
                        for lead in range(1,25):
                            target = release+timedelta(hours=lead)
                            forecast_keys.append((release.isoformat(),lead))
                            beyond_year += target.year > 2025
                meta['release_time_issues'] = issues
                meta['forecast_batches'] = len(rows)-1
                meta['duplicate_release_lead_keys'] = len(forecast_keys)-len(set(forecast_keys))
                meta['forecast_target_points_in_2026'] = beyond_year
                meta['regions'] = {'pv_forecast': numeric_region(rows,2,len(rows),3,26,'kW')}
            else:
                payload, placeholders, formulas = [], [], []
                for r,row in enumerate(rows,1):
                    for c,v in enumerate(row,1):
                        if isinstance(v,str) and v.startswith('='):
                            formulas.append(cell_address(r,c))
                        if isinstance(v,str) and any(x in v for x in ('…','⋮','⁝')):
                            placeholders.append(cell_address(r,c))
                if ws.title in {'计划购电量','调整购电量'}:
                    start_col = 2
                    for r in range(2,len(rows)+1):
                        for c in range(start_col,len(rows[r-1])+1):
                            if rows[r-1][c-1] not in (None,''):
                                payload.append(cell_address(r,c))
                    if path.name != 'result1.xlsx':
                        meta['date_axis'] = axis_dates([(r,rows[r-1][0]) for r in range(2,len(rows)+1)],date(2025,2,1),date(2025,12,31))
                    meta['first_period'] = plain(rows[1][0] if path.name=='result1.xlsx' else rows[0][1])
                    meta['last_period'] = plain(rows[-1][0] if path.name=='result1.xlsx' else rows[0][144])
                elif ws.title == '充放电量':
                    cols = [2,3,5] if path.name=='result1.xlsx' else [3,4,6]
                    for r in range(2,len(rows)+1):
                        for c in cols:
                            v=rows[r-1][c-1]
                            if v not in (None,'') and cell_address(r,c) not in placeholders:
                                payload.append(cell_address(r,c))
                else:
                    for r in range(2,len(rows)+1):
                        for c in [2,3]:
                            v=rows[r-1][c-1]
                            if v not in (None,'') and cell_address(r,c) not in placeholders:
                                payload.append(cell_address(r,c))
                meta.update({'unexpected_filled_result_cells':payload,'formula_cells':formulas,
                             'placeholder_cells':placeholders})
                result['templates'][key]=meta
                continue
            result['data'][key]=meta
        wb.close()
    result['actual_load_pv_price_axes_identical'] = len(set(json.dumps(v) for v in matrices.values())) == 1
    after = {p.relative_to(ROOT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
    if before != after:
        raise RuntimeError('Source changed during audit; no report saved')
    result['raw_hashes_unchanged'] = True
    if OUT.exists():
        raise FileExistsError(f'Preserve prior audit: {OUT}')
    OUT.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    brief = {'output':str(OUT),'raw_hashes_unchanged':True,'axes_identical':result['actual_load_pv_price_axes_identical'],
             'regions':{k:v['regions'] for k,v in result['data'].items()},
             'template_payload_issues':{k:v['unexpected_filled_result_cells'] for k,v in result['templates'].items() if v['unexpected_filled_result_cells']}}
    print(json.dumps(brief,ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
