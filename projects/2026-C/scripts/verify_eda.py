"""Independent raw workbook and standard-library checks for the EDA headline claims."""
from pathlib import Path
from collections import defaultdict
from datetime import datetime
from statistics import fmean, correlation
import csv
import json
import math
import openpyxl

P=Path(__file__).resolve().parents[1]
E=P/'audit/eda-v1'


def date_value(v):
    return v if isinstance(v,datetime) else datetime.fromisoformat(str(v))


def main():
    summary=json.loads((E/'summary.json').read_text(encoding='utf-8'))
    wb=openpyxl.load_workbook(P/'data/raw/C题/附件/附件2.xlsx',read_only=True,data_only=True)
    arrays={}
    monthly=defaultdict(list)
    weekdays=defaultdict(list)
    for sheet,col in [('小区负载','load_kw'),('光伏发电实际功率','pv_actual_kw')]:
        arrays[col]=[]
        for row in wb[sheet].iter_rows(min_row=2,values_only=True):
            day=date_value(row[0])
            vals=list(map(float,row[1:145]))
            assert len(vals)==144
            arrays[col].extend(vals)
            monthly[(day.month,col)].extend(vals)
            if col=='load_kw':
                weekdays[day.weekday()].extend(vals)
    wb.close()
    wb=openpyxl.load_workbook(P/'data/raw/C题/附件/附件4.xlsx',read_only=True,data_only=True)
    arrays['price_yuan_per_kwh']=[]
    for row in wb.active.iter_rows(min_row=2,values_only=True):
        day=date_value(row[0]); vals=list(map(float,row[1:145]))
        arrays['price_yuan_per_kwh'].extend(vals)
        monthly[(day.month,'price_yuan_per_kwh')].extend(vals)
    wb.close()
    arrays['net_load_kw']=[a-b for a,b in zip(arrays['load_kw'],arrays['pv_actual_kw'])]
    checks=[]
    def check(name,a,b,tol=1e-8):
        good=math.isclose(a,b,rel_tol=1e-10,abs_tol=tol)
        checks.append({'name':name,'independent':a,'recorded':b,'ok':good})
        assert good, name
    for col,v in arrays.items():
        check(col+'_mean',fmean(v),summary['annual'][col]['mean'])
        for h in [24,168]:
            check(f'{col}_lag_{h}',correlation(v[:-h*6],v[h*6:]),summary['lag_correlations'][col][str(h)]['raw'])
    for r in summary['monthly']:
        for col in ['load_kw','pv_actual_kw','price_yuan_per_kwh']:
            check(f'month_{r["month"]}_{col}',fmean(monthly[(r['month'],col)]),r[col])
    with (E/'weekday_summary.csv').open(encoding='utf-8',newline='') as f:
        for r in csv.DictReader(f):
            check('weekday_'+r['weekday'],fmean(weekdays[int(r['weekday'])]),float(r['load_kw']))
    check('surplus_fraction',sum(v<0 for v in arrays['net_load_kw'])/52560,summary['surplus']['sample_fraction'])
    # Verify same-target paired forecast errors via a dictionary, not the pandas filter code.
    groups=defaultdict(list)
    with (E/'forecast_matched.csv').open(encoding='utf-8',newline='') as f:
        for r in csv.DictReader(f):
            if r['pv_actual_kw']:
                groups[r['target_datetime']].append(r)
    pairs=[]
    for rows in groups.values():
        for r in rows:
            if int(r['release_hour'])==0 and int(r['lead_hours'])>=7:
                future=[x for x in rows if r['release_datetime']<x['release_datetime']<r['target_datetime']]
                if future:
                    latest=max(future,key=lambda x:x['release_datetime'])
                    actual=float(r['pv_actual_kw'])
                    pairs.append((actual,abs(float(r['pv_forecast_kw'])-actual),abs(float(latest['pv_forecast_kw'])-actual)))
    for scope,rows in [('all',pairs),('actual_positive',[x for x in pairs if x[0]>0])]:
        s=summary['forecast']['same_target'][scope]
        check(scope+'_n',len(rows),s['n'])
        check(scope+'_initial_mae',fmean(x[1] for x in rows),s['initial_mae_kw'])
        check(scope+'_latest_mae',fmean(x[2] for x in rows),s['latest_mae_kw'])
        check(scope+'_worse_fraction',sum(x[2]>x[1] for x in rows)/len(rows),s['worse_fraction'])
    # Annual same-clock means compared to supplied representative reference (descriptive only).
    with (E/'reference_profile.csv').open(encoding='utf-8',newline='') as f:
        refs=list(csv.DictReader(f))
    agreement={}
    for col,rc in [('load_kw','load_kw'),('pv_actual_kw','pv_forecast_kw'),('price_yuan_per_kwh','price_yuan_per_kwh')]:
        diffs=[float(refs[i][rc])-fmean(arrays[col][i::144]) for i in range(144)]
        agreement[col]={'max_abs':max(map(abs,diffs)),'mae':fmean(map(abs,diffs))}
    rec={'raw_workbooks_read':['附件2.xlsx','附件4.xlsx'],'checks':checks,'all_ok':all(x['ok'] for x in checks),'weekday_date_note':'2025-01-01 is Wednesday; source calendar preserved, no weekday shift.','reference_vs_annual_same_label':agreement,'not_proven':'Neither data-generation provenance nor causal explanations or future forecasting accuracy.'}
    (E/'verification.json').write_text(json.dumps(rec,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'checks':len(checks),'all_ok':rec['all_ok'],'reference_agreement':agreement},ensure_ascii=False))


if __name__=='__main__':
    main()
