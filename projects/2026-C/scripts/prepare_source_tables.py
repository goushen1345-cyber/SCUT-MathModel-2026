"""Execute approved Data-plan-v1. Copy source values; no interval interpolation.

python -X utf8 projects/2026-C/scripts/prepare_source_tables.py
Use --output <new project-relative directory> for non-overwriting reproduction.
"""
import argparse
import csv
from datetime import date, datetime, time, timedelta
import hashlib
import json
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
NS = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
SOURCES = {
    'annex1': ('data/raw/C题/附件/附件1.xlsx', 'Sheet1'),
    'annex2_load': ('data/raw/C题/附件/附件2.xlsx', '小区负载'),
    'annex2_pv': ('data/raw/C题/附件/附件2.xlsx', '光伏发电实际功率'),
    'annex3': ('data/raw/C题/附件/附件3.xlsx', 'Sheet1'),
    'annex4': ('data/raw/C题/附件/附件4.xlsx', 'Sheet1'),
}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def label(value):
    return value.isoformat() if isinstance(value,(date,datetime,time)) else str(value or '')


def parse_date(value):
    if isinstance(value,datetime):
        return value.date()
    if isinstance(value,date):
        return value
    return date(*map(int,re.split(r'[-/.]',value.strip())))


def minutes(value):
    if isinstance(value,time):
        assert value.second == 0 and value.microsecond == 0
        return value.hour*60+value.minute
    m=re.fullmatch(r'(\d{1,2}):(\d{1,2})(\+1)?',value.strip())
    assert m, value
    return int(m[1])*60+int(m[2])+(1440 if m[3] else 0)


def source(sid):
    path,sheet=SOURCES[sid]
    wb=openpyxl.load_workbook(ROOT/path,read_only=True,data_only=False)
    ws=wb[sheet]
    rows=list(ws.iter_rows(values_only=True))
    wb.close()
    with zipfile.ZipFile(ROOT/path) as z:
        tree=ET.fromstring(z.read('xl/workbook.xml'))
        rid=next(s for s in tree.find('s:sheets',NS) if s.attrib['name']==sheet).attrib[REL]
        target=next(r.attrib['Target'] for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels')) if r.attrib['Id']==rid)
        target=target.lstrip('/') if target.startswith('/') else 'xl/'+target
        values={}
        for cell in ET.fromstring(z.read(target)).findall('.//s:sheetData/s:row/s:c',NS):
            v=cell.find('s:v',NS)
            if v is not None and cell.attrib.get('t','n')=='n':
                values[cell.attrib['r']]=v.text
    return rows,values


def write_csv(out,name,fields,records):
    path=out/name
    with path.open('x',encoding='utf-8-sig',newline='') as stream:
        writer=csv.DictWriter(stream,fieldnames=fields)
        writer.writeheader()
        writer.writerows(records)
    return {'rows':len(records),'columns':fields,'sha256':sha(path)}


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--output',default='data/prepared-v1')
    args=parser.parse_args()
    out=(ROOT/args.output).resolve()
    if ROOT not in out.parents or (ROOT/'data/raw').resolve() in out.parents or out==(ROOT/'data/raw').resolve():
        raise ValueError('Output must be a new project directory outside raw')
    if out.exists():
        raise FileExistsError(f'Will not overwrite: {out}')
    audit=json.loads((ROOT/'audit/data_audit_v1.json').read_text(encoding='utf-8'))
    expected=audit['source_sha256']
    if any(sha(ROOT/p)!=h for p,h in expected.items()):
        raise RuntimeError('Source hashes differ from approved audit')
    loaded={sid:source(sid) for sid in SOURCES}
    tables={}
    rows,raw=loaded['annex1']
    records=[]
    for r in range(2,len(rows)+1):
        t=rows[r-1][0]
        records.append(dict(slot_index=r-1,source_time_label=label(t),source_time_type=type(t).__name__,source_minute=minutes(t),price_yuan_per_kwh=raw[f'B{r}'],load_kw=raw[f'C{r}'],pv_forecast_kw=raw[f'D{r}'],source_id='annex1',time_cell=f'A{r}',price_cell=f'B{r}',load_cell=f'C{r}',pv_forecast_cell=f'D{r}'))
    tables['day_reference.csv']=records

    def matrix(sid):
        rows,raw=loaded[sid]
        result={}
        for r in range(2,len(rows)+1):
            d=parse_date(rows[r-1][0])
            for c in range(2,len(rows[0])+1):
                t=rows[0][c-1]
                key=(d.isoformat(),minutes(t))
                if key in result:
                    raise ValueError(f'Duplicate source key: {sid} {key}')
                result[key]=dict(source_date=d.isoformat(),source_date_label=label(rows[r-1][0]),source_date_type=type(rows[r-1][0]).__name__,source_time_label=label(t),source_time_type=type(t).__name__,source_minute=minutes(t),value=raw[f'{get_column_letter(c)}{r}'],source_id=sid,date_cell=f'A{r}',time_cell=f'{get_column_letter(c)}1',value_cell=f'{get_column_letter(c)}{r}')
        return result
    load,pv,price=(matrix(s) for s in ['annex2_load','annex2_pv','annex4'])
    if set(load)!=set(pv) or set(load)!=set(price):
        raise ValueError('Actual axes do not align exactly')
    common=['source_date','source_date_label','source_date_type','source_time_label','source_time_type','source_minute']
    records=[]
    for key in sorted(load):
        a,b=load[key],pv[key]
        assert all(a[k]==b[k] for k in common)
        record={k:a[k] for k in common}
        record.update(load_kw=a['value'],pv_actual_kw=b['value'])
        for prefix,item in [('load',a),('pv',b)]:
            record.update({f'{prefix}_{k}':item[k] for k in ['source_id','date_cell','time_cell','value_cell']})
        records.append(record)
    tables['actual_power.csv']=records
    records=[]
    for key in sorted(price):
        a=price[key]
        record={k:a[k] for k in common}
        record.update(price_yuan_per_kwh=a['value'])
        record.update({k:a[k] for k in ['source_id','date_cell','time_cell','value_cell']})
        records.append(record)
    tables['actual_price.csv']=records

    actual_labels={datetime.fromisoformat(d)+timedelta(minutes=m) for d,m in load}
    rows,raw=loaded['annex3']
    records=[]
    inherited_rows=0
    for group_row in range(2,len(rows)+1,4):
        group_date=rows[group_row-1][0]
        d=parse_date(group_date)
        for offset in range(4):
            r=group_row+offset
            release_value=rows[r-1][1]
            release=datetime.combine(d,time())+timedelta(minutes=minutes(release_value))
            inherited=int(rows[r-1][0] in (None,''))
            inherited_rows+=inherited
            for lead in range(1,25):
                c=lead+2
                target=release+timedelta(hours=lead)
                records.append(dict(release_datetime=release.isoformat(),target_datetime=target.isoformat(),lead_hours=lead,pv_forecast_kw=raw[f'{get_column_letter(c)}{r}'],source_date_label=label(group_date),source_date_type=type(group_date).__name__,source_row_date_label=label(rows[r-1][0]),source_row_date_type=type(rows[r-1][0]).__name__,source_release_label=label(release_value),source_release_type=type(release_value).__name__,date_inherited=inherited,has_actual_target_label=int(target in actual_labels),source_id='annex3',group_date_cell=f'A{group_row}',row_date_cell=f'A{r}',release_cell=f'B{r}',value_cell=f'{get_column_letter(c)}{r}'))
    tables['pv_forecasts.csv']=records
    expected_rows={'day_reference.csv':144,'actual_power.csv':52560,'actual_price.csv':52560,'pv_forecasts.csv':35040}
    assert {k:len(v) for k,v in tables.items()}==expected_rows
    out.mkdir(parents=True,exist_ok=False)
    outputs={name:write_csv(out,name,list(records[0]),records) for name,records in tables.items()}
    readme='''# 整理后的来源数据（prepared-v1）

按已批准的最小格式整理方案生成，当前是待用户确认的候选版本。
保留全部原始数值、来源位置和发布版本，没有训练、插值、能量换算或填写购电结果。

| 文件 | 每行表示什么 | 行数 | 数值单位 |
|---|---|---:|---|
| day_reference.csv | 附件1代表日的一个原始时间标签 | 144 | 功率kW；电价元/kWh |
| actual_power.csv | 一个来源日期和原时间标签的实际负载及光伏 | 52560 | kW |
| actual_price.csv | 一个来源日期和原时间标签的实际价格 | 52560 | 元/kWh |
| pv_forecasts.csv | 一次发布针对一个未来整点的光伏预报 | 35040 | kW |

## 时间字段

source_date保留原表归属日；代表日没有日历日期，不假造日期。
source_minute为距离归属日零点的分钟数：0:10对应10，0:00+1对应1440。
source_time_label保留原单元格值的文字表示：Excel时间转ISO文字，原文本不改；source_time_type保留原类型。
这只是来源标签，不说明功率代表前一段、后一段或平均值，不能直接把它当作已确认的购电区间。
CSV保留日期和时间为明确文字、分钟数为整数。不要覆盖CSV后再声称原值未变。

## 预报字段

release_datetime是发布时刻；lead_hours为1—24；target_datetime等于发布时刻加提前小时数。
同一目标时刻的不同发布版本全部保留，不能按目标时刻去重或用较晚发布值覆盖较早发布。
date_inherited=1表示原行日期空白，来自同组首行；原行日期及组首单元格都保留。
has_actual_target_label只表示目标标签是否在附件2中存在，不代表当时可提前获知其实际值。
全年最后实际标签是2025-12-31的0:00+1，即2026-01-01 00:00；36条预报晚于此时刻，仅标注而未删除。

## 信息身份与单位

actual_power是实际观测；actual_price是实际价格，在Q4中不可提前知道未来部分。
day_reference来自附件1，其负载/预测用于Q1和已确认的Q2首日启动参考；其电价用于固定日曲线场景。
pv_forecasts来自附件3，仅在相应发布时刻起可用，不能用于未经授权的Q2/Q4-2预报输入。
CSV本身不强制这些权限，后续研究代码还需落实时间切片。
功率仍为kW，价格仍为元/kWh；没有得到区间能量kWh，未填补数值、裁剪、平滑、标准化或插值。

## 数值及来源

数值列直接保存XLSX内部的数值文本，没有按显示小数位四舍五入；重读时按数值解析。
source_id在manifest.json的sources中映射到原文件与工作表。
所有*_cell列均为原工作表单元格地址；实际功率的load_/pv_前缀分别指两张原表。
source_date_label/type保留原日期值表示及类型；source_row_date_label/type保留预报当前行的原日期空白。
source_release_label/type保留预报发布时间的原值表示及类型。
slot_index仅是附件1原行顺序索引，不是新定义的物理区间。

manifest.json记录文件哈希、来源身份、行数、字段、生成脚本和授权。
逐值核对结果另存项目audit/prepared_v1_verification.json，不是人工核验或数据冻结。

## 复现

在仓库根目录使用Python运行：
python -X utf8 projects/2026-C/scripts/prepare_source_tables.py --output tmp/prepared-v1-reproduction
只允许新目录；脚本拒绝覆盖已有输出或写入原始材料目录。
核验命令：python -X utf8 projects/2026-C/scripts/verify_prepared_tables.py
'''
    (out/'README.md').write_text(readme,encoding='utf-8')
    if any(sha(ROOT/p)!=h for p,h in expected.items()):
        raise RuntimeError('Original files changed during preparation')
    manifest={'version':'prepared-v1','status':'candidate_pending_user_confirmation','plan_version':'Data-plan-v1','authorization':'DATA-PREP-012','created_at':datetime.now().astimezone().isoformat(),'python':sys.version,'openpyxl':openpyxl.__version__,'script':'scripts/prepare_source_tables.py','script_sha256':sha(Path(__file__)),'raw_sha256':expected,'raw_unchanged_after_preparation':True,'sources':{sid:{'path':path,'sheet':sheet} for sid,(path,sheet) in SOURCES.items()},'outputs':outputs,'readme_sha256':sha(out/'README.md'),'transformations':{'original_numeric_values_changed':0,'original_numeric_cells':193152,'total_output_records':sum(expected_rows.values()),'forecast_rows_with_inherited_group_date':inherited_rows,'forecast_records_without_actual_target_label':sum(not r['has_actual_target_label'] for r in tables['pv_forecasts.csv']),'numeric_text':'copied verbatim from original XLSX numeric value elements','physical_interval_mapping':'not selected','hourly_interpolation':'not performed','power_to_energy':'not performed'},'information_roles':{'day_reference.csv':'attachment1_reference_not_Q2_future_actuals','actual_power.csv':'actual_observation_not_future_information','actual_price.csv':'Q4_actual_price_not_known_ahead','pv_forecasts.csv':'available_only_at_or_after_release_in_authorized_branches'}}
    (out/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'output':str(out),'rows':expected_rows,'numeric_values_changed':0,'inherited_date_rows':inherited_rows,'forecast_outside_actual_labels':manifest['transformations']['forecast_records_without_actual_target_label']},ensure_ascii=False))


if __name__=='__main__':
    main()
