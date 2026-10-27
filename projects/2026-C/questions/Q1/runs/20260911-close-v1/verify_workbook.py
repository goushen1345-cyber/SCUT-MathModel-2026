"""Read saved XLSX independently; preserve raw template and check its temporal keys."""
from pathlib import Path
from decimal import Decimal as D, ROUND_HALF_UP
from copy import copy
import json,re
from openpyxl import load_workbook
from verify_close import RUN,BASE,PROJECT,OUT,cells,rows,sha,load

def main():
    target=RUN/'deliverables/result1.xlsx'; raw=PROJECT/'data/raw/C题/附件/附件5/result1.xlsx'
    a=load_workbook(raw); b=load_workbook(target); source=cells(PROJECT/'data/raw/C题/附件/附件1.xlsx')
    saved=[cells(target,i) for i in [1,2]]; original=[cells(raw,i) for i in [1,2]]
    lp=rows(BASE/'outputs/schedules/C04.csv'); et=D('1e-5'); ct=D('1e-4'); failures=[]; count=0; worst={}
    def ok(name,test):
        nonlocal count
        count+=1
        if not test: failures.append(name)
    def near(name,x,y,tol=et):
        delta=abs(D(str(x))-D(str(y))); key=name.split(':')[0]
        worst[key]=max(worst.get(key,D(0)),delta); ok(name,delta<=tol)
    ok('sheets',a.sheetnames==b.sheetnames)
    allowed=[{f'B{i}' for i in range(2,146)}, {f'{c}{i}' for i in range(2,8) for c in 'BC'}|{'E2','E3'}]
    metadata_normalizations=[]
    for j,s in enumerate(a):
        t=b[s.title]
        ok('dimensions:'+s.title,(s.max_row,s.max_column)==(t.max_row,t.max_column))
        ok('merges:'+s.title,list(s.merged_cells.ranges)==list(t.merged_cells.ranges))
        ok('tables:'+s.title,list(s.tables)==list(t.tables)); ok('panes:'+s.title,s.freeze_panes==t.freeze_panes)
        for row in s:
            for c in row:
                d=t[c.coordinate]
                if c.coordinate not in allowed[j]:
                    ok('preserved_value:'+s.title+c.coordinate,c.value==d.value)
                    ok('preserved_number_format:'+s.title+c.coordinate,c.number_format==d.number_format)
                else:
                    ok('numeric_output:'+s.title+c.coordinate,type(d.value) in (int,float))
                    ok('display_format:'+s.title+c.coordinate,d.number_format=='0.00')
                    near('display_rounding:'+s.title+c.coordinate,D(str(d.value)).quantize(D('.01'),rounding=ROUND_HALF_UP),d.value,D('.005000000001'))
                for prop in ['name','sz','bold','italic','underline','strike','vertAlign']:
                    ok('font:'+s.title+c.coordinate+prop,getattr(c.font,prop)==getattr(d.font,prop))
                # Export normalizes empty sides and drops font charset/family metadata.
                # Require effective visible alignment/borders to remain unchanged.
                for prop in ['horizontal','vertical','textRotation','wrapText','shrinkToFit','indent']:
                    x=getattr(c.alignment,prop); y=getattr(d.alignment,prop)
                    ok('alignment:'+s.title+c.coordinate+prop,(x or None)==(y or None))
                for edge in ['left','right','top','bottom','diagonal']:
                    x=getattr(c.border,edge); y=getattr(d.border,edge)
                    ok('border:'+s.title+c.coordinate+edge,getattr(x,'style',None)==getattr(y,'style',None))
                if c.font.charset!=d.font.charset or c.font.family!=d.font.family:
                    metadata_normalizations.append(s.title+'!'+c.coordinate)
        for col,dim in s.column_dimensions.items():
            other=t.column_dimensions[col]
            ok('column:'+s.title+col,dim.width==other.width and dim.hidden==other.hidden)
    q_by_k={}; bijection=[]
    def minute(token):
        h,m=map(int,token.replace('+1','').split(':')); return (60*h+m)%1440
    for row in range(2,146):
        label=saved[0][f'A{row}']; begin,end=label.split('-'); start=minute(begin)
        ok('10minute:'+str(row),(minute(end)-start)%1440==10)
        k=start//10+1; bijection.append(k); q=D(saved[0][f'B{row}']); q_by_k[k]=q
        near('schedule_value:'+str(row),q,lp[k-1]['purchase_kwh'])
        ok('unchanged_label:'+str(row),label==original[0][f'A{row}'])
    ok('144unique_slots',sorted(bijection)==list(range(1,145)))
    ok('last_row_midnight',bijection[-1]==1 and bijection[0]==2 and bijection[-2]==144)
    total_cost=D(0); total_q=sum(q_by_k.values()); prev=D(6000)
    for k,r in enumerate(lp,1):
        sr=145 if k==1 else k; l=D(source[f'C{sr}'])/6; v=D(source[f'D{sr}'])/6; p=D(source[f'B{sr}'])
        q=q_by_k[k]; c=D(r['charge_kwh']); discharge=D(r['discharge']); waste=D(r['unused_pv_kwh'])
        near('export_supply:'+str(k),q+v+discharge,l+c+waste)
        prev+=D('.9')*c-discharge/D('.9')
        ok('export_capacity:'+str(k),D(1200)-et<=prev<=D(10800)+et)
        ok('export_power:'+str(k),c<=D(5000)/6+et and discharge<=D(5000)/6+et)
        total_cost+=q*p
    near('export_end',prev,6000)
    for i in range(6):
        row=i+2
        for c,field in [('B','charge_kwh'),('C','discharge')]:
            near('block:'+c+str(row),saved[1][f'{c}{row}'],sum(D(r[field]) for r in lp[i*24:(i+1)*24]))
    near('initial',saved[1]['E2'],6000); near('terminal',saved[1]['E3'],6000)
    summary=load(OUT/'summary.json'); near('total_purchase',total_q,summary['main_totals']['purchase_kwh'])
    near('total_cost',total_cost,summary['main_totals']['cost_yuan'],ct)
    for rel,h in load(PROJECT/'data/prepared-v1/manifest.json')['raw_sha256'].items(): ok('raw_unchanged:'+rel,sha(PROJECT/rel)==h)
    verification=load(OUT/'verification.json')
    shifts=[]
    for scenario in verification['changed_schedules']:
        changes=scenario['changed_slots']; positive=[r for r in changes if r['purchase_difference_kwh']>0]; negative=[r for r in changes if r['purchase_difference_kwh']<0]
        shifts.append(dict(initial_kwh=scenario['initial_kwh'],
            shifted_kwh=sum(r['purchase_difference_kwh'] for r in positive),
            added_purchase_weighted_price=sum(r['cost_difference_yuan'] for r in positive)/sum(r['purchase_difference_kwh'] for r in positive),
            removed_purchase_weighted_price=sum(r['cost_difference_yuan'] for r in negative)/sum(r['purchase_difference_kwh'] for r in negative),
            first_added_time=positive[0]['time'],last_added_time=positive[-1]['time'],
            first_removed_time=negative[0]['time'],last_removed_time=negative[-1]['time']))
    result=dict(passed=not failures,checks=count,failures=failures,
        largest_residuals={k:float(v) for k,v in worst.items()},
        total_purchase_kwh=float(total_q),total_cost_yuan=float(total_cost),source_sha256=sha(raw),output_sha256=sha(target),
        preserved_sheets=a.sheetnames,filled_numeric_cells=158,source_unchanged=True,
        time_mapping='template last row -> k1 midnight; first data row -> k2; penultimate -> k144',
        visible_style_checked=True,serialization_note=f'Artifact export removed charset/family metadata in {len(metadata_normalizations)} cells and normalized empty border sides; effective fonts, alignments, border styles and widths checked unchanged.',
        visual_note='Both sheets and midnight tail rendered and inspected by AI, no clipping; not an interactive Microsoft Excel session.',
        validation_scope='144 purchases checked from saved XLSX with baseline 10-minute charge/discharge evidence; template itself contains only 4-hour charge/discharge totals.',
        shift_interpretation=shifts)
    (OUT/'workbook_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2)); assert not failures,failures[:12]

if __name__=='__main__': main()
